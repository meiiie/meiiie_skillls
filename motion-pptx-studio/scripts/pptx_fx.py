#!/usr/bin/env python3
"""v6.18 native craft helpers (python-pptx), all playing in PowerPoint 2010+ / LibreOffice WITHOUT Morph.
Techniques and when to use them: references/techniques.md. Helpers, not templates.

  from pptx_fx import text_fill_image, mask_picture, parallax_effects, push_in_effects
  text_fill_image(shape, 'library/ch1/hero.jpg')   # every run of a text box filled with the photo (image inside a word)
  mask_picture(pic, 'ellipse'|'roundRect'|'hexagon'|'triangle'|'flowChartDelay'|...)   # crop a picture to a shape
  parallax_effects(['!!bg', '!!mid', '!!fg'], dx_in=1.2, ms=6000)  -> motion.json "effects" list: three layers travel
      by 1:2.5:6 of dx along motion paths at once (depth parallax with classic paths, not Morph)
  push_in_effects('!!hero', scale=1.12, ms=8000) -> slow Ken-Burns grow (classic grow/shrink emphasis)
CLI self-test: python3 pptx_fx.py selftest IMAGE OUT.pptx"""
import copy, sys
from lxml import etree
from pptx.oxml.ns import qn

def _blip_fill(part, image):
    _, rid = part.get_or_add_image_part(image)
    bf = etree.SubElement(etree.Element('x'), qn('a:blipFill'), rotWithShape='1')
    etree.SubElement(bf, qn('a:blip')).set(qn('r:embed'), rid)
    etree.SubElement(etree.SubElement(bf, qn('a:stretch')), qn('a:fillRect'))
    return bf

def text_fill_image(shape, image):
    """Fill the glyphs of every run with a photo. Use on >= 200 pt heavy type so the photo reads."""
    part = shape.part
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            for tag in ('a:solidFill', 'a:gradFill', 'a:blipFill', 'a:noFill', 'a:pattFill'):
                for e in rPr.findall(qn(tag)): rPr.remove(e)
            bf = _blip_fill(part, image)
            ln = rPr.find(qn('a:ln')); (ln.addnext(bf) if ln is not None else rPr.insert(0, bf))
    return shape

def mask_picture(pic, prst='ellipse'):
    """Crop a picture to a preset geometry (circle, arch-like flowChartDelay, hexagon, roundRect ...)."""
    spPr = pic._element.spPr
    g = spPr.find(qn('a:prstGeom'))
    if g is None: g = etree.SubElement(spPr, qn('a:prstGeom'))
    g.set('prst', prst)
    if g.find(qn('a:avLst')) is None: etree.SubElement(g, qn('a:avLst'))
    return pic

def parallax_effects(names, dx_in=1.2, dy_in=0.0, ms=6000, ratios=(1, 2.5, 6)):
    """Depth parallax for motion.json: far layer moves least. Shapes must be oversize so edges never show."""
    return [{'shape': n, 'fx': 'path', 'path': [[0, 0], [-dx_in * k / ratios[-1], -dy_in * k / ratios[-1]]],
             'ms': ms, 'click': 0, 'delay': 0} for i, (n, k) in enumerate(zip(names, ratios))]

def push_in_effects(name, scale=1.12, ms=8000):
    return [{'shape': name, 'fx': 'grow', 'scale': scale, 'ms': ms, 'click': 0, 'delay': 0}]

if __name__ == '__main__' and sys.argv[1:2] == ['selftest']:
    from pptx import Presentation; from pptx.util import Inches, Pt
    img, out = sys.argv[2], sys.argv[3]
    pr = Presentation(); pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)
    s = pr.slides.add_slide(pr.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.3), Inches(0.5), Inches(12.7), Inches(5)); r = tb.text_frame.paragraphs[0].add_run()
    r.text = 'SAN HÔ'; r.font.size = Pt(260); r.font.bold = True
    text_fill_image(tb, img)
    mask_picture(s.shapes.add_picture(img, Inches(9.5), Inches(4.2), Inches(3), Inches(3)), 'ellipse')
    pr.save(out); print('wrote', out)
