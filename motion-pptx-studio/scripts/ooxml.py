"""Small, dependency-light toolkit for editing a .pptx as raw OOXML.

Subcommands:
  unpack DECK.pptx DIR              unzip a deck into DIR (no reformatting)
  pack DIR OUT.pptx                 zip DIR back ([Content_Types].xml first); checks every XML parses
  order DIR                         print deck position -> slide file (and the !! layers on it)
  inventory DIR [--slides a.xml,b.xml]
                                    list shapes: id, name, kind, x/y/w/h in inches, text preview
  sequence DIR SRC.xml [SRC.xml ...]
                                    set the deck order. A file listed again is DUPLICATED (with its
                                    notes) so one layout can become several Morph keyframes.
                                    Slides not listed leave the slide list (files stay until `clean`).
  clean DIR                         delete slide files not in the slide list, their notes/rels,
                                    content-type overrides and media nothing points to
  ensure-notes DECK.pptx            give every slide a notes page (python-pptx); run before unpack
  notes-get DIR [--out notes.md]    export speaker notes as '##### N' blocks (deck positions)
  notes-set DIR notes.md            write '##### N' blocks into the notes pages (deck positions)

Positions vs files: position N is the N-th slide shown; slideK.xml is the file. After
`sequence` they differ. motion.json durations use FILE numbers; notes use POSITIONS.
"""
import argparse
import copy
import re
import shutil
import sys
import zipfile
from pathlib import Path

from lxml import etree

NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
}
P, A, R = ('{%s}' % NS[k] for k in ('p', 'a', 'r'))
REL = '{%s}' % NS['rel']
CT = '{%s}' % NS['ct']
EMU = 914400
SLIDE_T = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide'
NOTES_T = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide'
SLIDE_CT = 'application/vnd.openxmlformats-officedocument.presentationml.slide+xml'
NOTES_CT = 'application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml'


def parse(path):
    return etree.parse(str(path))


def write(tree, path):
    tree.write(str(path), xml_declaration=True, encoding='UTF-8', standalone=True)


def num(name):
    return int(re.findall(r'\d+', Path(name).name)[0])


# ---------- zip ----------
def unpack(deck, out):
    out = Path(out)
    if out.exists() and any(out.iterdir()):
        raise SystemExit(f'{out} is not empty; unpack into a fresh folder')
    out.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(deck) as z:
        z.extractall(out)
    print('unpacked', deck, '->', out)


def pack(src, out):
    src = Path(src)
    bad = []
    files = [p for p in src.rglob('*') if p.is_file() and p.name not in ('.DS_Store', 'Thumbs.db')]
    for p in files:
        if p.suffix in ('.xml', '.rels'):
            try:
                etree.parse(str(p))
            except etree.XMLSyntaxError as e:
                bad.append(f'{p.relative_to(src)}: {e}')
    if bad:
        raise SystemExit('XML does not parse, not packing:\n' + '\n'.join(bad))
    ct = src / '[Content_Types].xml'
    if not ct.is_file():
        raise SystemExit('missing [Content_Types].xml in ' + str(src))
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(ct, '[Content_Types].xml')
        for p in sorted(files):
            rel = p.relative_to(src).as_posix()
            if rel != '[Content_Types].xml':
                z.write(p, rel)
    print('packed', out, f'{Path(out).stat().st_size / 1e6:.2f} MB')


# ---------- presentation model ----------
class Deck:
    def __init__(self, d):
        self.d = Path(d)
        self.ppt = self.d / 'ppt'
        self.pres_path = self.ppt / 'presentation.xml'
        self.rels_path = self.ppt / '_rels' / 'presentation.xml.rels'
        self.ct_path = self.d / '[Content_Types].xml'
        if not self.pres_path.is_file():
            raise SystemExit('not an unpacked deck (no ppt/presentation.xml): ' + str(d))
        self.pres = parse(self.pres_path)
        self.rels = parse(self.rels_path)
        self.ct = parse(self.ct_path)

    def rid_to_file(self):
        out = {}
        for r in self.rels.getroot():
            if r.get('Type') == SLIDE_T:
                out[r.get('Id')] = Path(r.get('Target')).name
        return out

    def order(self):
        m = self.rid_to_file()
        lst = self.pres.getroot().find('p:sldIdLst', NS)
        return [] if lst is None else [m[s.get(R + 'id')] for s in lst]

    def save(self):
        write(self.pres, self.pres_path)
        write(self.rels, self.rels_path)
        write(self.ct, self.ct_path)

    def slide_rels(self, f):
        return self.ppt / 'slides' / '_rels' / f'{f}.rels'

    def notes_file(self, f):
        rp = self.slide_rels(f)
        if not rp.is_file():
            return None
        for r in parse(rp).getroot():
            if r.get('Type') == NOTES_T:
                return Path(r.get('Target')).name
        return None


def shape_rows(slide_path):
    t = parse(slide_path)
    rows = []
    for c in t.iter(P + 'cNvPr'):
        el = c.getparent().getparent()
        kind = etree.QName(el).localname
        xf = el.find('./p:spPr/a:xfrm', NS)
        if xf is None:
            xf = el.find('./p:grpSpPr/a:xfrm', NS)
        if xf is None:
            xf = el.find('./p:xfrm', NS)
        geo = ''
        if xf is not None and xf.find('a:off', NS) is not None:
            o, e = xf.find('a:off', NS), xf.find('a:ext', NS)
            geo = '%.2f,%.2f %.2fx%.2f' % (int(o.get('x')) / EMU, int(o.get('y')) / EMU, int(e.get('cx')) / EMU, int(e.get('cy')) / EMU)
        txt = ' / '.join(''.join(t_.text or '' for t_ in p.iter(A + 't')) for p in el.iter(A + 'p')) if kind == 'sp' else ''
        rows.append((c.get('id'), c.get('name'), kind, geo, re.sub(r'\s+', ' ', txt).strip()[:70]))
    return rows


def cmd_order(d):
    deck = Deck(d)
    for i, f in enumerate(deck.order(), 1):
        names = sorted({r[1] for r in shape_rows(deck.ppt / 'slides' / f) if r[1] and r[1].startswith('!!')})
        print(f'{i:3}  {f:14} {" ".join(names)}')


def cmd_inventory(d, slides):
    deck = Deck(d)
    files = slides.split(',') if slides else deck.order()
    for f in files:
        print(f'== {f}')
        for sid, name, kind, geo, txt in shape_rows(deck.ppt / 'slides' / f):
            print(f'  {sid:>4} {kind:8} {name[:34]:34} {geo:24} {txt}')


def cmd_sequence(d, order):
    deck = Deck(d)
    rid_of = {v: k for k, v in deck.rid_to_file().items()}
    for f in order:
        if f not in rid_of:
            raise SystemExit(f'{f} is not a slide of this deck')
    slides_dir = deck.ppt / 'slides'
    notes_dir = deck.ppt / 'notesSlides'
    next_slide = max(num(p) for p in slides_dir.glob('slide*.xml')) + 1
    next_notes = max([num(p) for p in notes_dir.glob('notesSlide*.xml')] or [0]) + 1
    rels_root = deck.rels.getroot()
    next_rid = max(int(re.sub(r'\D', '', r.get('Id')) or 0) for r in rels_root) + 1
    pres_root = deck.pres.getroot()
    lst = pres_root.find('p:sldIdLst', NS)
    ids = {s.get(R + 'id'): int(s.get('id')) for s in lst}
    next_id = max(ids.values()) + 1
    ct_root = deck.ct.getroot()
    seen, entries, printed = set(), [], []
    for f in order:
        if f not in seen:
            seen.add(f)
            entries.append((ids[rid_of[f]], rid_of[f]))
            printed.append(f)
            continue
        new = f'slide{next_slide}.xml'
        next_slide += 1
        shutil.copy(slides_dir / f, slides_dir / new)
        rels = parse(deck.slide_rels(f))
        for r in rels.getroot():
            if r.get('Type') == NOTES_T:
                old_notes = Path(r.get('Target')).name
                new_notes = f'notesSlide{next_notes}.xml'
                next_notes += 1
                shutil.copy(notes_dir / old_notes, notes_dir / new_notes)
                nrels = parse(notes_dir / '_rels' / f'{old_notes}.rels')
                for nr in nrels.getroot():
                    if Path(nr.get('Target')).name == f:
                        nr.set('Target', '../slides/' + new)
                write(nrels, notes_dir / '_rels' / f'{new_notes}.rels')
                r.set('Target', '../notesSlides/' + new_notes)
                etree.SubElement(ct_root, CT + 'Override', PartName=f'/ppt/notesSlides/{new_notes}', ContentType=NOTES_CT)
        write(rels, deck.slide_rels(new))
        etree.SubElement(ct_root, CT + 'Override', PartName=f'/ppt/slides/{new}', ContentType=SLIDE_CT)
        rid = f'rId{next_rid}'
        next_rid += 1
        etree.SubElement(rels_root, REL + 'Relationship', Id=rid, Type=SLIDE_T, Target='slides/' + new)
        entries.append((next_id, rid))
        next_id += 1
        printed.append(new)
    for s in list(lst):
        lst.remove(s)
    for sid, rid in entries:
        e = etree.SubElement(lst, P + 'sldId', id=str(sid))
        e.set(R + 'id', rid)
    deck.save()
    for i, f in enumerate(printed, 1):
        print(f'position {i}: {f}')


def cmd_clean(d):
    deck = Deck(d)
    keep = set(deck.order())
    slides_dir = deck.ppt / 'slides'
    ct_root = deck.ct.getroot()
    rels_root = deck.rels.getroot()
    removed = []
    for p in slides_dir.glob('slide*.xml'):
        if p.name in keep:
            continue
        nf = deck.notes_file(p.name)
        if nf:
            for q in (deck.ppt / 'notesSlides' / nf, deck.ppt / 'notesSlides' / '_rels' / f'{nf}.rels'):
                q.unlink(missing_ok=True)
            for o in ct_root.findall(CT + 'Override'):
                if o.get('PartName') == f'/ppt/notesSlides/{nf}':
                    ct_root.remove(o)
        deck.slide_rels(p.name).unlink(missing_ok=True)
        p.unlink()
        for o in ct_root.findall(CT + 'Override'):
            if o.get('PartName') == f'/ppt/slides/{p.name}':
                ct_root.remove(o)
        for r in rels_root.findall(REL + 'Relationship'):
            if r.get('Type') == SLIDE_T and Path(r.get('Target')).name == p.name:
                rels_root.remove(r)
        removed.append(p.name)
    deck.save()
    # orphan media: not referenced by any remaining .rels
    refs = set()
    for rp in deck.d.rglob('*.rels'):
        for r in parse(rp).getroot():
            refs.add(Path(r.get('Target', '')).name)
    media = deck.ppt / 'media'
    orphans = [m for m in media.iterdir()] if media.is_dir() else []
    orphans = [m for m in orphans if m.name not in refs]
    for m in orphans:
        m.unlink()
    print(f'removed {len(removed)} unused slides, {len(orphans)} orphan media files')


def cmd_ensure_notes(deck_path):
    from pptx import Presentation
    prs = Presentation(deck_path)
    made = 0
    for s in prs.slides:
        if not s.has_notes_slide:
            s.notes_slide  # creates it
            made += 1
    prs.save(deck_path)
    print(f'notes pages created: {made} (of {len(prs.slides)})')


def notes_body(tree):
    for sp in tree.iter(P + 'sp'):
        ph = sp.find('.//p:nvPr/p:ph', NS)
        if ph is not None and ph.get('type') == 'body':
            return sp.find('p:txBody', NS)
    return None


def cmd_notes_get(d, out):
    deck = Deck(d)
    blocks = []
    for i, f in enumerate(deck.order(), 1):
        nf = deck.notes_file(f)
        text = ''
        if nf:
            body = notes_body(parse(deck.ppt / 'notesSlides' / nf))
            if body is not None:
                text = '\n'.join(''.join(t.text or '' for t in p.iter(A + 't')) for p in body.findall('a:p', NS))
        blocks.append(f'##### {i} {f}\n{text.strip()}\n')
    s = '\n'.join(blocks)
    if out:
        Path(out).write_text(s, encoding='utf-8')
        print('wrote', out)
    else:
        sys.stdout.write(s)


def cmd_notes_set(d, src, lang):
    deck = Deck(d)
    order = deck.order()
    text = Path(src).read_text(encoding='utf-8')
    done = []
    for blk in text.split('##### ')[1:]:
        head, _, body_text = blk.partition('\n')
        pos = int(head.split()[0])
        if not 1 <= pos <= len(order):
            raise SystemExit(f'block ##### {pos}: deck has {len(order)} slides')
        f = order[pos - 1]
        nf = deck.notes_file(f)
        if not nf:
            raise SystemExit(f'position {pos} ({f}) has no notes page: pack, run `ensure-notes`, unpack again')
        path = deck.ppt / 'notesSlides' / nf
        tree = parse(path)
        body = notes_body(tree)
        if body is None:
            raise SystemExit(f'{nf} has no body placeholder')
        template = body.find('a:p', NS)
        for p in body.findall('a:p', NS):
            body.remove(p)
        for line in body_text.rstrip().split('\n') or ['']:
            p = etree.SubElement(body, A + 'p')
            r = etree.SubElement(p, A + 'r')
            rp = etree.SubElement(r, A + 'rPr', lang=lang, dirty='0')
            etree.SubElement(r, A + 't').text = line
        write(tree, path)
        done.append(pos)
    print('notes written for positions', done)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    s = sub.add_parser('unpack'); s.add_argument('deck'); s.add_argument('dir')
    s = sub.add_parser('pack'); s.add_argument('dir'); s.add_argument('out')
    s = sub.add_parser('order'); s.add_argument('dir')
    s = sub.add_parser('inventory'); s.add_argument('dir'); s.add_argument('--slides')
    s = sub.add_parser('sequence'); s.add_argument('dir'); s.add_argument('files', nargs='+')
    s = sub.add_parser('clean'); s.add_argument('dir')
    s = sub.add_parser('ensure-notes'); s.add_argument('deck')
    s = sub.add_parser('notes-get'); s.add_argument('dir'); s.add_argument('--out')
    s = sub.add_parser('notes-set'); s.add_argument('dir'); s.add_argument('notes'); s.add_argument('--lang', default='vi-VN')
    a = ap.parse_args(argv)
    if a.cmd == 'unpack': unpack(a.deck, a.dir)
    elif a.cmd == 'pack': pack(a.dir, a.out)
    elif a.cmd == 'order': cmd_order(a.dir)
    elif a.cmd == 'inventory': cmd_inventory(a.dir, a.slides)
    elif a.cmd == 'sequence': cmd_sequence(a.dir, a.files)
    elif a.cmd == 'clean': cmd_clean(a.dir)
    elif a.cmd == 'ensure-notes': cmd_ensure_notes(a.deck)
    elif a.cmd == 'notes-get': cmd_notes_get(a.dir, a.out)
    elif a.cmd == 'notes-set': cmd_notes_set(a.dir, a.notes, a.lang)


if __name__ == '__main__':
    main()
