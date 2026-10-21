"""Preview the deck's motion as an MP4 WITHOUT PowerPoint: a labelled SIMULATION.

  python3 preview_motion.py DECK.pptx --out WORK/preview/motion-sim.mp4 [--fps 30] [--width 1280]
                           [--hold 1.2] [--click-gap 0.7] [--max-seconds 90 [--cut]] [--slides 1-8] [--keep DIR]
                           [--lang vi] [--label TEXT] [--label-size PX]
  --max-seconds fits the whole deck into that length (holds shrink first, then transitions); --cut stops at
  that second instead. --lang vi tags every frame "MÔ PHỎNG · không phải PowerPoint chạy thật".

How it works (LibreOffice renders only static slides, so the motion is reconstructed):
  1. every top-level shape of every slide is rendered ALONE by LibreOffice on black and on white
     (difference matte -> exact alpha), plus the background of each slide;
  2. the slide's real p:transition and p:timing XML (what motion.py wrote) are parsed: transition
     type/duration, every effect's preset, start, duration, direction, iterate (letters/words),
     motion path, scale/rotation, repeat and auto-reverse, click groups, triggers, embedded video;
  3. frames are composited: Morph = same-name shapes interpolate position/size and crossfade,
     unmatched ones fade; push/wipe/split/cover/uncover/dissolve/curtains/fall over/... are 2-D
     approximations; entrances, emphasis, exits and paths are drawn per shape; video shapes play
     their MP4; clicks and triggers are 'pressed' automatically after a short gap.
Every frame carries a "SIMULATION" tag. Real PowerPoint playback differs in detail (3-D transitions,
letter timing, Morph of text and geometry, z-order while a shape animates). For the real thing,
export from PowerPoint (File > Export > Create a Video) on a machine that has it.
"""
import argparse
import copy
import io
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parent))
import motionlib  # noqa: E402

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RNS = NS['r']
P = '{%s}' % NS['p']
SHAPE_TAGS = {'sp', 'pic', 'grpSp', 'cxnSp', 'graphicFrame'}
PAGE = 1.6       # sprite page is 1.6x the slide so shadows and glows are not clipped
MARGIN = 0.25    # inches of margin captured around each shape


# ------------------------------------------------------------------ 1. sprites
def xfrm_of(el):
    tag = etree.QName(el).localname
    if tag == 'grpSp':
        return el.find('p:grpSpPr/a:xfrm', NS)
    if tag == 'graphicFrame':
        return el.find('p:xfrm', NS)
    return el.find('p:spPr/a:xfrm', NS)


def build_sprite_decks(src, tmp, sel):
    from pptx import Presentation
    meta = {}
    for variant, col in (('black', '000000'), ('white', 'FFFFFF')):
        prs = Presentation(src)
        W, H = prs.slide_width, prs.slide_height
        orig = list(prs.slides)
        n0 = len(orig)
        prs.slide_width, prs.slide_height = int(W * PAGE), int(H * PAGE)
        page = 0
        for si in sel:
            s = orig[si - 1]
            recs = []
            for z, el in enumerate([e for e in s.shapes._spTree if etree.QName(e).localname in SHAPE_TAGS]):
                xf = xfrm_of(el)
                c = el.find('.//p:cNvPr', NS)
                rec = {'z': z, 'id': c.get('id'), 'name': c.get('name') or '', 'tag': etree.QName(el).localname}
                if xf is None or xf.find('a:off', NS) is None:
                    continue
                ns = prs.slides.add_slide(s.slide_layout)
                for ph in list(ns.shapes):
                    ph._element.getparent().remove(ph._element)
                cSld = ns._element.find('p:cSld', NS)
                cSld.insert(0, etree.fromstring(f'<p:bg xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"><p:bgPr><a:solidFill>'
                                                f'<a:srgbClr val="{col}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'))
                e = copy.deepcopy(el)
                xf2 = xfrm_of(e)
                off, ext = xf2.find('a:off', NS), xf2.find('a:ext', NS)
                x, y, w, h = int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))
                rot = int(xf2.get('rot', '0')) / 60000.0
                if 'rot' in xf2.attrib:
                    del xf2.attrib['rot']
                dx, dy = int(W * PAGE / 2 - w / 2) - x, int(H * PAGE / 2 - h / 2) - y
                off.set('x', str(x + dx))
                off.set('y', str(y + dy))
                # (a group's children follow its a:off through chOff/chExt, so moving off moves them)
                for node in e.iter():
                    for k, v in list(node.attrib.items()):
                        if k.startswith('{%s}' % RNS) and v in s.part.rels:
                            rel = s.part.rels[v]
                            nid = (ns.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                                   else ns.part.relate_to(rel.target_part, rel.reltype))
                            node.set(k, nid)
                ns.shapes._spTree.append(e)
                rec.update(x=x / W, y=y / H, w=w / W, h=h / H, rot=rot, page=page,
                           text=''.join(t.text or '' for t in el.iter('{%s}t' % NS['a'])))
                vf = el.find('.//a:videoFile', NS)
                if vf is not None:
                    rid = vf.get('{%s}link' % RNS)
                    rel = s.part.rels.get(rid)
                    if rel is not None and not rel.is_external:
                        rec['video'] = rel.target_part.partname.lstrip('/')
                recs.append(rec)
                page += 1
            if variant == 'black':
                meta[si] = recs
        lst = prs.slides._sldIdLst
        for sid in list(lst)[:n0]:
            prs.part.drop_rel(sid.get('{%s}id' % RNS))
            lst.remove(sid)
        prs.save(tmp / f'sprites_{variant}.pptx')
    # backgrounds: the original slides with every shape removed
    prs = Presentation(src)
    keep = set(sel)
    for i, s in enumerate(prs.slides, 1):
        for el in [e for e in s.shapes._spTree if etree.QName(e).localname in SHAPE_TAGS]:
            el.getparent().remove(el)
    lst = prs.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(lst), 1))):
        if i not in keep:
            prs.part.drop_rel(sid.get('{%s}id' % RNS))
            lst.remove(sid)
    prs.save(tmp / 'base.pptx')
    return meta


def to_pngs(pptx, outdir, width, profile):
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'pdf', '--outdir', str(outdir), str(pptx)],
                   capture_output=True, timeout=1800)
    pdf = outdir / (pptx.stem + '.pdf')
    if not pdf.exists():
        raise SystemExit(f'LibreOffice could not render {pptx.name}')
    subprocess.run(['pdftoppm', '-png', '-scale-to-x', str(width), '-scale-to-y', '-1', str(pdf), str(outdir / 'p')], check=True)
    return sorted(outdir.glob('p-*.png'), key=lambda p: int(re.findall(r'(\d+)\.png$', p.name)[0]))


class Sprite:
    def __init__(self, rec, black, white, W, H):
        self.rec = rec
        self.name = rec['name']
        self.id = rec['id']
        sw, sh = black.size
        mx, my = MARGIN / 13.333 * W, MARGIN / 7.5 * H
        w, h = rec['w'] * W, rec['h'] * H
        cx, cy = sw / 2, sh / 2
        box = (int(cx - w / 2 - mx), int(cy - h / 2 - my), int(math.ceil(cx + w / 2 + mx)), int(math.ceil(cy + h / 2 + my)))
        # a shape larger than the page (a Morph zoom plate at x3) renders only inside the page: crop to it, so a
        # 30-inch circle does not ask PIL for a 190-megapixel crop (DecompressionBombError, v5.2 fix)
        box = (max(0, box[0]), max(0, box[1]), min(sw, box[2]), min(sh, box[3]))
        b = np.asarray(black.crop(box).convert('RGB'), np.float32)
        wt = np.asarray(white.crop(box).convert('RGB'), np.float32)
        alpha = np.clip(1 - (wt - b).mean(axis=2) / 255.0, 0, 1)
        col = np.where(alpha[..., None] > 1e-3, b / np.maximum(alpha[..., None], 1e-3), 0)
        rgba = np.dstack([np.clip(col, 0, 255), alpha * 255]).astype(np.uint8)
        self.img = Image.fromarray(rgba, 'RGBA')
        bb = self.img.getbbox()
        self.empty = bb is None
        # rest position of the crop's top-left on the slide (px)
        self.x0 = rec['x'] * W - mx - ((cx - w / 2 - mx) - box[0])
        self.y0 = rec['y'] * H - my - ((cy - h / 2 - my) - box[1])
        self.rot = rec.get('rot', 0)
        if self.rot:
            self.img = self.img.rotate(-self.rot, resample=Image.BICUBIC, expand=False)
        self.video = rec.get('video')

    @property
    def center(self):
        return self.x0 + self.img.width / 2, self.y0 + self.img.height / 2


# ------------------------------------------------------------------ 2. timing
def parse_path(d):
    toks = re.findall(r'[MLCZE]|-?[\d.]+(?:e-?\d+)?', d)
    pts, cur, i, cmd = [], (0.0, 0.0), 0, None
    while i < len(toks):
        t = toks[i]
        if t in 'MLCZE':
            cmd = t
            i += 1
            if t in 'ZE':
                continue
        if cmd == 'C':
            c1 = (float(toks[i]), float(toks[i + 1])); c2 = (float(toks[i + 2]), float(toks[i + 3])); p = (float(toks[i + 4]), float(toks[i + 5]))
            for k in range(1, 13):
                u = k / 12
                pts.append(tuple((1 - u) ** 3 * cur[j] + 3 * (1 - u) ** 2 * u * c1[j] + 3 * (1 - u) * u ** 2 * c2[j] + u ** 3 * p[j] for j in range(2)))
            cur = p
            i += 6
        else:
            cur = (float(toks[i]), float(toks[i + 1]))
            pts.append(cur)
            i += 2
    return pts or [(0, 0)]


def along(pts, f):
    if len(pts) == 1:
        return pts[0]
    seg = [math.dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1)]
    total = sum(seg) or 1e-9
    d = f * total
    for i, s in enumerate(seg):
        if d <= s or i == len(seg) - 1:
            u = 0 if s == 0 else min(1, d / s)
            return (pts[i][0] + (pts[i + 1][0] - pts[i][0]) * u, pts[i][1] + (pts[i + 1][1] - pts[i][1]) * u)
        d -= s
    return pts[-1]


def delay_of(ctn):
    for c in ctn.findall('p:stCondLst/p:cond', NS):
        d = c.get('delay')
        if d and d != 'indefinite':
            try:
                return int(d)
            except ValueError:
                return 0
    return 0


def effect_info(ctn):
    """One effect's essentials from its preset cTn."""
    e = {'cls': ctn.get('presetClass'), 'id': int(ctn.get('presetID', '0')), 'sub': int(ctn.get('presetSubtype', '0')),
         'loop': ctn.get('repeatCount') == 'indefinite', 'kids': []}
    tgt = ctn.find('.//p:spTgt', NS)
    e['spid'] = tgt.get('spid') if tgt is not None else None
    it = ctn.find('p:iterate', NS)
    if it is not None:
        tp = it.find('p:tmPct', NS)
        e['iterate'] = (it.get('type', 'el'), (int(tp.get('val')) / 100000.0) if tp is not None else 0.1)
    durs = []
    for b in ctn.find('p:childTnLst', NS) if ctn.find('p:childTnLst', NS) is not None else []:
        tag = etree.QName(b).localname
        c = b.find('.//p:cTn', NS)
        dur = c.get('dur') if c is not None else None
        dur = int(dur) if dur and dur.isdigit() else 0
        start = delay_of(c) if c is not None else 0
        k = {'tag': tag, 'dur': dur, 'start': start, 'autoRev': c is not None and c.get('autoRev') == '1'}
        if tag == 'animEffect':
            k['filter'] = b.get('filter', '')
            k['out'] = b.get('transition') == 'out'
        elif tag == 'animMotion':
            k['path'] = parse_path(b.get('path', 'M 0 0 E'))
        elif tag == 'animScale':
            by = b.find('p:by', NS)
            k['by'] = (int(by.get('x')) / 100000.0) if by is not None else 1.0
        elif tag == 'animRot':
            k['by'] = int(b.get('by', '0')) / 60000.0
        elif tag == 'anim':
            names = [a.text for a in b.iter(P + 'attrName')]
            k['attr'] = names[0] if names else ''
            k['vals'] = [tv.find('.//p:strVal', NS).get('val') if tv.find('.//p:strVal', NS) is not None else None
                         for tv in b.iter(P + 'tav')]
        elif tag == 'set':
            v = b.find('.//p:strVal', NS)
            k['val'] = v.get('val') if v is not None else None
        e['kids'].append(k)
        durs.append(start + dur * (2 if k['autoRev'] else 1))
    e['dur'] = max(durs) if durs else 1
    if e['cls'] == 'mediacall':
        e['dur'] = max(durs) if durs else 1
    return e


def parse_timing(root, click_gap_ms):
    """Effects with absolute start (ms after the transition) for the auto-play preview."""
    out = []
    timing = root.find(P + 'timing')
    if timing is None:
        return out, 0
    t_cursor = 0

    def walk(ctn, base):
        lst = []
        for par in ctn.findall('p:childTnLst/p:par', NS):
            c = par.find('p:cTn', NS)
            st = base + delay_of(c)
            if c.get('presetClass'):
                e = effect_info(c)
                e['start'] = st
                lst.append(e)
            else:
                lst += walk(c, st)
        return lst

    for seq in timing.iter(P + 'seq'):
        c = seq.find('p:cTn', NS)
        if c.get('nodeType') == 'mainSeq':
            for g in c.findall('p:childTnLst/p:par', NS):
                gc = g.find('p:cTn', NS)
                auto = any(cc.get('evt') == 'onBegin' for cc in gc.findall('p:stCondLst/p:cond', NS))
                start = t_cursor if auto or not out else t_cursor + click_gap_ms
                effs = walk(gc, start)
                # afterEffect: a group starts when the previous one ended
                for e in effs:
                    e['click'] = not auto
                out += effs
                if effs:
                    t_cursor = max(e['start'] + (e['dur'] if not e['loop'] else min(e['dur'], 1500)) for e in effs)
    for seq in timing.iter(P + 'seq'):
        c = seq.find('p:cTn', NS)
        if c.get('nodeType') == 'interactiveSeq':
            effs = walk(c, t_cursor + click_gap_ms)
            for e in effs:
                e['trigger'] = True
            out += effs
            if effs:
                t_cursor = max(e['start'] + e['dur'] for e in effs)
    return out, t_cursor


# ------------------------------------------------------------------ 3. drawing
def ease_out(u):
    u = max(0.0, min(1.0, u))
    return 1 - (1 - u) ** 3


def ease_io(u):
    u = max(0.0, min(1.0, u))
    return 4 * u ** 3 if u < 0.5 else 1 - (-2 * u + 2) ** 3 / 2


def reveal_mask(size, frac, mode, soft=0.04):
    """Alpha multiplier for wipe/split style reveals; mode: from-left/right/top/bottom, split-h, split-v."""
    w, h = size
    frac = max(0.0, min(1.0, frac))
    if mode in ('l', 'r'):
        x = np.linspace(0, 1, w, dtype=np.float32)[None, :]
        if mode == 'r':
            x = 1 - x
        m = np.clip((frac * (1 + soft) - x) / soft, 0, 1)
        return np.repeat(m, h, axis=0)
    if mode in ('t', 'b'):
        y = np.linspace(0, 1, h, dtype=np.float32)[:, None]
        if mode == 'b':
            y = 1 - y
        m = np.clip((frac * (1 + soft) - y) / soft, 0, 1)
        return np.repeat(m, w, axis=1)
    if mode == 'sv':  # split opening from a vertical centre line
        x = np.abs(np.linspace(-1, 1, w, dtype=np.float32))[None, :]
        return np.repeat(np.clip((frac * (1 + soft) - x) / soft, 0, 1), h, axis=0)
    if mode == 'sh':
        y = np.abs(np.linspace(-1, 1, h, dtype=np.float32))[:, None]
        return np.repeat(np.clip((frac * (1 + soft) - y) / soft, 0, 1), w, axis=1)
    return np.ones((h, w), np.float32)


FLY_DIR = {1: 't', 2: 'r', 4: 'b', 8: 'l'}


class State:
    __slots__ = ('visible', 'alpha', 'dx', 'dy', 'scale', 'rot', 'mask', 'video_t', 'stretch')

    def __init__(self):
        self.visible, self.alpha, self.dx, self.dy, self.scale, self.rot, self.mask, self.video_t = True, 1.0, 0.0, 0.0, 1.0, 0.0, None, None
        self.stretch = None   # (edge, fraction): v5.5 Stretch entrance, the shape grows out of that edge


STRETCH_SUB = {1: 't', 2: 'r', 4: 'b', 8: 'l', 10: 'h', 5: 'v'}


def sweep_mask(size, frac, soft=0.03):
    """Wheel (1 spoke): a clockwise sweep from 12 o'clock, i.e. Manim's Create on a circle."""
    w, h = size
    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32)
    ang = (np.arctan2(xx - w / 2, -(yy - h / 2)) / (2 * np.pi)) % 1.0
    return np.clip((frac * (1 + soft) - ang) / soft, 0, 1)


def shape_state(sp, effs, t, W, H):
    """Visual state of one sprite at time t (ms after the slide's transition)."""
    st = State()
    entr = [e for e in effs if e['cls'] == 'entr']
    if entr and t < min(e['start'] for e in entr):
        st.visible = False
    for e in sorted(effs, key=lambda e: e['start']):
        u_ms = t - e['start']
        if u_ms < 0:
            continue
        dur = max(1, e['dur'])
        cls, pid, sub = e['cls'], e['id'], e['sub']
        if cls == 'entr':
            st.visible = True
            # letters/words: the shape's effect lasts dur*(1 + (n-1)*pct); reveal left to right in steps
            if e.get('iterate'):
                n = max(1, len(sp.rec.get('text', '')) if e['iterate'][0] == 'lt' else len(sp.rec.get('text', '').split()))
                base = max(1, min(k['dur'] for k in e['kids'] if k['dur']) if any(k['dur'] for k in e['kids']) else 500)
                step = base * e['iterate'][1]
                total = base + step * (n - 1)
                if u_ms < total:
                    k = min(n, int(u_ms / max(1, step)) + 1)
                    frac = k / n
                    m = reveal_mask(sp.img.size, frac, 'l', soft=0.6 / n)
                    st.mask = m if st.mask is None else st.mask * m
                continue
            u = u_ms / dur
            if u >= 1:
                continue
            if pid == 1:
                pass
            elif pid == 10:
                st.alpha *= ease_out(u)
            elif pid == 2:  # fly / drop
                d = FLY_DIR.get(sub, 'b')
                k = 1 - ease_out(u)
                cx, cy = sp.center
                if d == 'b':
                    st.dy += k * (H + sp.img.height / 2 - cy)
                elif d == 't':
                    st.dy -= k * (cy + sp.img.height / 2)
                elif d == 'l':
                    st.dx -= k * (cx + sp.img.width / 2)
                else:
                    st.dx += k * (W + sp.img.width / 2 - cx)
            elif pid == 42:
                st.alpha *= ease_out(u)
                st.dy += (1 - ease_out(u)) * 0.06 * H
            elif pid == 22:
                m = reveal_mask(sp.img.size, ease_out(u), FLY_DIR.get(sub, 'l'))
                st.mask = m if st.mask is None else st.mask * m
            elif pid == 16:
                m = reveal_mask(sp.img.size, ease_out(u), 'sv' if sub in (21, 37) else 'sh')
                st.mask = m if st.mask is None else st.mask * m
            elif pid == 53:
                st.scale *= max(0.02, ease_out(u))
                st.alpha *= ease_out(u)
            elif pid == 21:   # v5.5: a real clockwise sweep, not a fade
                m = sweep_mask(sp.img.size, ease_out(u))
                st.mask = m if st.mask is None else st.mask * m
            elif pid == 17:   # v5.5 Stretch: grows out of one edge
                st.stretch = (STRETCH_SUB.get(sub, 'b'), max(0.002, ease_out(u)))
            else:
                st.alpha *= ease_out(u)
        elif cls == 'exit':
            u = u_ms / dur
            if u >= 1:
                st.visible = False
                continue
            if pid == 2:
                d = FLY_DIR.get(sub, 'b')
                k = ease_io(u)
                st.dy += k * H if d == 'b' else (-k * H if d == 't' else 0)
                st.dx += -k * W if d == 'l' else (k * W if d == 'r' else 0)
            elif pid == 53:
                st.scale *= 1 - ease_io(u)
            elif pid == 22:
                m = 1 - reveal_mask(sp.img.size, ease_io(u), FLY_DIR.get(sub, 'l'))
                st.mask = m if st.mask is None else st.mask * m
            else:
                st.alpha *= 1 - u
        elif cls == 'emph':
            for k in e['kids']:
                kd = max(1, k['dur'])
                period = kd * (2 if k['autoRev'] else 1)
                v = u_ms - k['start']
                if v < 0:
                    continue
                if e['loop']:
                    v = v % period
                elif v >= period:
                    v = period if not k['autoRev'] else 0
                ph = v / kd
                amt = ease_io(ph if ph <= 1 else 2 - ph)
                if k['tag'] == 'animScale':
                    st.scale *= 1 + (k['by'] - 1) * amt
                elif k['tag'] == 'animRot':
                    st.rot += k['by'] * (ease_io(min(1, v / kd)))
                elif k['tag'] == 'anim' and k.get('attr') == 'style.visibility':
                    st.visible = st.visible and ((v / kd) % 1) >= 0.5
        elif cls == 'path':
            for k in e['kids']:
                if k['tag'] != 'animMotion':
                    continue
                kd = max(1, k['dur'])
                v = u_ms - k['start']
                if e['loop']:
                    v = v % kd
                f = ease_io(min(1, v / kd)) if not e['loop'] else min(1, v / kd)
                px, py = along(k['path'], f)
                st.dx += px * W
                st.dy += py * H
        elif cls == 'mediacall':
            st.video_t = u_ms / 1000.0
    return st


class VideoReader:
    def __init__(self, data, size, fps):
        self.tmp = Path(tempfile.mkdtemp(prefix='pv-vid-'))
        self.path = self.tmp / 'v.mp4'
        self.path.write_bytes(data)
        self.size, self.fps = size, fps
        self.proc = subprocess.Popen(['ffmpeg', '-v', 'error', '-i', str(self.path), '-vf', f'scale={size[0]}:{size[1]},fps={fps}',
                                      '-f', 'rawvideo', '-pix_fmt', 'rgba', '-'], stdout=subprocess.PIPE)
        self.idx, self.last = -1, None

    def frame(self, t):
        want = int(t * self.fps)
        n = self.size[0] * self.size[1] * 4
        while self.idx < want:
            buf = self.proc.stdout.read(n)
            if len(buf) < n:
                break
            self.last = Image.frombuffer('RGBA', self.size, buf, 'raw', 'RGBA', 0, 1).copy()
            self.idx += 1
        return self.last


def paste(canvas, img, x, y, alpha=1.0, scale=1.0, rot=0.0, mask=None, stretch=None):
    if alpha <= 0.003:
        return
    if mask is not None:
        a = np.asarray(img.getchannel('A'), np.float32) * mask
        img = img.copy()
        img.putalpha(Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), 'L'))
    if stretch is not None:   # anchored one-axis scale (Stretch entrance)
        edge, f = stretch
        w0, h0 = img.width, img.height
        if edge in ('t', 'b', 'v'):
            nh = max(1, int(round(h0 * f)))
            img = img.resize((w0, nh), Image.BILINEAR)
            y += {'b': h0 - nh, 't': 0}.get(edge, (h0 - nh) / 2)
        else:
            nw = max(1, int(round(w0 * f)))
            img = img.resize((nw, h0), Image.BILINEAR)
            x += {'r': w0 - nw, 'l': 0}.get(edge, (w0 - nw) / 2)
    cx, cy = x + img.width / 2, y + img.height / 2
    if abs(scale - 1) > 1e-3:
        nw, nh = max(1, int(img.width * scale)), max(1, int(img.height * scale))
        img = img.resize((nw, nh), Image.BILINEAR)
    if abs(rot) > 0.05:
        img = img.rotate(-rot, resample=Image.BICUBIC, expand=True)
    if alpha < 0.997:
        a = img.getchannel('A').point(lambda v: int(v * alpha))
        img = img.copy()
        img.putalpha(a)
    canvas.alpha_composite(img, (int(round(cx - img.width / 2)), int(round(cy - img.height / 2)))) if \
        (0 <= int(round(cx - img.width / 2)) and 0 <= int(round(cy - img.height / 2)) and
         int(round(cx - img.width / 2)) + img.width <= canvas.width and int(round(cy - img.height / 2)) + img.height <= canvas.height) else \
        _paste_clipped(canvas, img, int(round(cx - img.width / 2)), int(round(cy - img.height / 2)))


def _paste_clipped(canvas, img, x, y):
    x0, y0 = max(0, x), max(0, y)
    x1, y1 = min(canvas.width, x + img.width), min(canvas.height, y + img.height)
    if x1 <= x0 or y1 <= y0:
        return
    canvas.alpha_composite(img.crop((x0 - x, y0 - y, x1 - x, y1 - y)), (x0, y0))


class Slide:
    def __init__(self, n, root, base, sprites, click_gap, videos):
        self.n = n
        self.base = base
        self.sprites = [s for s in sprites if not s.empty or s.video]
        self.trans, ms = motionlib.describe_transition(root)
        self.trans_ms = int(ms) if ms and str(ms).isdigit() else 700
        tr = root.find('.//' + P + 'transition')
        self.trans_dir = None
        if tr is not None:
            kids = [k for k in tr if isinstance(k.tag, str)]
            if kids:
                self.trans_dir = kids[0].get('dir')
        if self.trans in ('none', 'cut-in'):
            self.trans_ms = 0
        if self.trans == 'cut':
            self.trans_ms = 0
        self.effects, self.end_ms = parse_timing(root, click_gap)
        self.by_spid = {}
        for e in self.effects:
            self.by_spid.setdefault(e['spid'], []).append(e)
        self.videos = videos

    def frame(self, t, W, H, skip_animated=False):
        cv = self.base.copy()
        for sp in self.sprites:
            effs = self.by_spid.get(sp.id, [])
            st = shape_state(sp, effs, t, W, H)
            if not st.visible:
                continue
            img = sp.img
            if sp.video and st.video_t is not None and sp.video in self.videos:
                fr = self.videos[sp.video].frame(st.video_t)
                if fr is not None:
                    vimg = Image.new('RGBA', img.size, (0, 0, 0, 0))
                    mx, my = int(MARGIN / 13.333 * W), int(MARGIN / 7.5 * H)
                    vimg.paste(fr.resize((img.width - 2 * mx, img.height - 2 * my)), (mx, my))
                    img = vimg
            paste(cv, img, sp.x0 + st.dx, sp.y0 + st.dy, st.alpha, st.scale, st.rot, st.mask, st.stretch)
        return cv

    def rest_rects(self):
        return {s.name: s for s in self.sprites if s.name}


def morph_frame(A, B, ta, k, W, H):
    """Morph by object: same-name sprites interpolate, others crossfade."""
    k = ease_io(k)
    cv = Image.blend(A.base.convert('RGB'), B.base.convert('RGB'), k).convert('RGBA')
    bn = {s.name: s for s in B.sprites if s.name}
    an = {s.name: s for s in A.sprites if s.name}
    for sp in A.sprites:
        st = shape_state(sp, A.by_spid.get(sp.id, []), ta, W, H)
        if not st.visible:
            continue
        if sp.name in bn:
            b = bn[sp.name]
            stb = shape_state(b, B.by_spid.get(b.id, []), 0, W, H)
            if not stb.visible:
                paste(cv, sp.img, sp.x0, sp.y0, (1 - k) * st.alpha)
                continue
            x = sp.x0 + (b.x0 - sp.x0) * k
            y = sp.y0 + (b.y0 - sp.y0) * k
            w = sp.img.width + (b.img.width - sp.img.width) * k
            h = sp.img.height + (b.img.height - sp.img.height) * k
            ia = sp.img.resize((max(1, int(w)), max(1, int(h))), Image.BILINEAR)
            ib = b.img.resize((max(1, int(w)), max(1, int(h))), Image.BILINEAR)
            paste(cv, ia, x, y, 1 - k)
            paste(cv, ib, x, y, k)
        else:
            paste(cv, sp.img, sp.x0 + st.dx, sp.y0 + st.dy, (1 - k) * st.alpha, st.scale, st.rot, st.mask)
    for b in B.sprites:
        if b.name in an:
            continue
        stb = shape_state(b, B.by_spid.get(b.id, []), 0, W, H)
        if stb.visible:
            paste(cv, b.img, b.x0, b.y0, k)
    return cv


def transition_frame(A, B, fa, fb, k, kind, d, W, H, rng):
    a, b = fa.convert('RGB'), fb.convert('RGB')
    e = ease_io(k)
    if kind in ('fade', 'cut-in') or kind is None:
        return Image.blend(a, b, e)
    if kind in ('push', 'pan', 'conveyor', 'ferris', 'gallery'):
        d = d or 'u'
        cv = Image.new('RGB', (W, H))
        if d in ('u', 'd'):
            s = int(e * H) * (1 if d == 'u' else -1)
            cv.paste(a, (0, -s)); cv.paste(b, (0, H - s if d == 'u' else -H - s))
        else:
            s = int(e * W) * (1 if d == 'l' else -1)
            cv.paste(a, (-s, 0)); cv.paste(b, (W - s if d == 'l' else -W - s, 0))
        return cv
    if kind in ('cover', 'switch', 'flythrough'):
        d = d or 'l'
        cv = a.copy()
        off = {'l': (int((1 - e) * W), 0), 'r': (-int((1 - e) * W), 0), 'u': (0, int((1 - e) * H)), 'd': (0, -int((1 - e) * H))}.get(d[:1], (int((1 - e) * W), 0))
        cv.paste(b, off)
        return cv
    if kind in ('uncover',):
        cv = b.copy()
        cv.paste(a, (-int(e * W), 0))
        return cv
    if kind in ('wipe', 'reveal', 'blinds', 'randombars', 'comb', 'checker'):
        m = reveal_mask((W, H), e, {'r': 'l', 'l': 'r', 'u': 'b', 'd': 't'}.get(d or 'r', 'l'), soft=0.08)
        if kind in ('blinds', 'randombars', 'comb', 'checker'):
            bands = 10
            y = (np.arange(H) % (H // bands)) / (H // bands)
            m = np.repeat((y[:, None] < e).astype(np.float32), W, axis=1)
        return Image.composite(b, a, Image.fromarray((m * 255).astype(np.uint8), 'L'))
    if kind in ('split', 'doors', 'window', 'shape', 'diamond', 'zoombox'):
        mode = 'sh' if (d == 'horz' or kind == 'split') else 'sv'
        if kind in ('shape', 'diamond', 'zoombox'):
            yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
            r = np.sqrt(((xx - W / 2) / W) ** 2 + ((yy - H / 2) / W) ** 2) / 0.6
            m = np.clip((e * 1.05 - r) / 0.03, 0, 1)
        else:
            m = reveal_mask((W, H), e, mode, soft=0.05)
        return Image.composite(b, a, Image.fromarray((m * 255).astype(np.uint8), 'L'))
    if kind in ('dissolve', 'glitter', 'shred', 'honeycomb', 'fracture', 'crush', 'vortex', 'ripple', 'flash'):
        cell = 24
        gw, gh = W // cell + 1, H // cell + 1
        noise = rng.random((gh, gw)).astype(np.float32) if not hasattr(transition_frame, '_n') else transition_frame._n
        transition_frame._n = noise
        m = (noise < e).astype(np.uint8) * 255
        m = Image.fromarray(m, 'L').resize((gw * cell, gh * cell), Image.NEAREST).crop((0, 0, W, H))
        out = Image.composite(b, a, m)
        if kind == 'flash':
            out = Image.blend(out, Image.new('RGB', (W, H), 'white'), max(0, 1 - abs(k - 0.5) * 4) * 0.8)
        return out
    if kind in ('curtains', 'drape', 'wind', 'prestige', 'peeloff', 'pagecurl', 'airplane', 'origami'):
        cv = b.copy()
        if kind in ('curtains', 'drape'):
            half = W // 2
            l = a.crop((0, 0, half, H)).resize((max(1, int(half * (1 - e))), int(H * (1 - 0.1 * e))))
            r = a.crop((half, 0, W, H)).resize((max(1, int(half * (1 - e))), int(H * (1 - 0.1 * e))))
            cv.paste(l, (0, 0)); cv.paste(r, (W - r.width, 0))
        else:  # the old slide lifts and slides away with a shadow
            s = 1 - 0.15 * e
            aa = a.resize((max(1, int(W * s)), max(1, int(H * s))))
            cv.paste(aa, (int(-e * W * 1.1 + (W - aa.width) / 2), int((H - aa.height) / 2 - e * H * 0.2)))
        return cv
    if kind in ('fallover', 'flip', 'cube', 'rotate', 'orbit', 'zoom'):
        cv = b.copy()
        if kind == 'zoom':
            s = 1 + 0.6 * e
            aa = a.resize((int(W * s), int(H * s))).crop((int(W * (s - 1) / 2), int(H * (s - 1) / 2), int(W * (s - 1) / 2) + W, int(H * (s - 1) / 2) + H))
            return Image.blend(aa, b, e)
        h2 = max(1, int(H * (1 - e) * (1 - 0.0)))
        aa = a.resize((W, h2))
        dark = Image.blend(aa, Image.new('RGB', aa.size, 'black'), 0.6 * e)
        cv.paste(dark, (0, H - h2))
        return cv
    return Image.blend(a, b, e)


LABELS = {'en': 'SIMULATION · not PowerPoint playback', 'vi': 'MÔ PHỎNG · không phải PowerPoint chạy thật'}
MIN_HOLD_MS = 900


def natural_hold(S, a):
    hold_ms = min(a.max_slide * 1000, max(a.min_slide * 1000, S.end_ms + a.hold * 1000))
    vids = [e for e in S.effects if e['cls'] == 'mediacall']
    if vids:
        hold_ms = min(a.max_slide * 1000, max(hold_ms, max(e['start'] + e['dur'] for e in vids) + 600))
    return hold_ms


def fit_budget(holds, trans, budget_ms):
    """(new holds in ms, transition factor) that fit every slide into budget_ms. Holds shrink first, each to no
    less than MIN_HOLD_MS (or its own length if shorter); then transitions shrink (to 15% at most).
    budget 0 = no limit."""
    holds = [float(h) for h in holds]
    total = sum(holds) + sum(trans)
    if not budget_ms or total <= budget_ms:
        return holds, 1.0
    floors = [min(h, MIN_HOLD_MS) for h in holds]
    room = budget_ms - sum(trans)
    if room <= sum(floors):
        return floors, max(0.15, (budget_ms - sum(floors)) / max(1, sum(trans)))
    lo, hi = 0.0, 1.0
    for _ in range(40):   # the scale k with sum(max(floor, h * k)) == room
        k = (lo + hi) / 2
        if sum(max(f, h * k) for f, h in zip(floors, holds)) > room:
            hi = k
        else:
            lo = k
    return [max(f, h * lo) for f, h in zip(floors, holds)], 1.0


def tag_frame(img, txt, font):
    d = ImageDraw.Draw(img)
    pad = 6
    tw = d.textlength(txt, font=font)
    x, y = img.width - tw - 2 * pad - 14, 12
    d.rounded_rectangle((x, y, x + tw + 2 * pad, y + font.size + 2 * pad), radius=6, fill=(0, 0, 0, 150))
    d.text((x + pad, y + pad - 1), txt, font=font, fill=(255, 255, 255, 230))
    return img


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck')
    ap.add_argument('--out', required=True)
    ap.add_argument('--fps', type=int, default=30)
    ap.add_argument('--width', type=int, default=1280)
    ap.add_argument('--hold', type=float, default=1.2, help='seconds to hold after the last effect')
    ap.add_argument('--min-slide', type=float, default=2.6, help='minimum seconds per slide')
    ap.add_argument('--max-slide', type=float, default=9.0, help='cap per slide (long videos and loops are cut)')
    ap.add_argument('--click-gap', type=float, default=0.7, help='seconds before a simulated click')
    ap.add_argument('--slides', default='', help='e.g. 1-8 or 3,4,5')
    ap.add_argument('--keep', help='keep the sprite renders here')
    ap.add_argument('--label', default=None, help='tag burnt into every frame (default from --lang)')
    ap.add_argument('--lang', default='en', help='vi = Vietnamese default label "MÔ PHỎNG · không phải PowerPoint chạy thật"')
    ap.add_argument('--label-size', type=int, default=0, help='label font px (default: width / 60)')
    ap.add_argument('--max-seconds', type=float, default=0,
                    help='fit the WHOLE deck into this many seconds by shortening holds (then transitions); 0 = natural length')
    ap.add_argument('--cut', action='store_true', help='with --max-seconds: stop at that second instead (the pre-5.4 behaviour)')
    a = ap.parse_args(argv)
    if a.label is None:
        a.label = LABELS.get(a.lang.split('-')[0].lower(), LABELS['en'])
    src = Path(a.deck).resolve()
    z = zipfile.ZipFile(src)
    pres = etree.fromstring(z.read('ppt/presentation.xml'))
    rels = etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
    rmap = {r.get('Id'): r.get('Target') for r in rels}
    parts = ['ppt/' + rmap[s.get('{%s}id' % RNS)] for s in pres.iter(P + 'sldId')]
    sel = list(range(1, len(parts) + 1))
    if a.slides:
        sel = []
        for chunk in a.slides.split(','):
            if '-' in chunk:
                lo, hi = map(int, chunk.split('-'))
                sel += list(range(lo, hi + 1))
            else:
                sel.append(int(chunk))
    tmp = Path(a.keep or tempfile.mkdtemp(prefix='pv-'))
    tmp.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix='pv-lo-'))
    W = a.width // 2 * 2
    H = int(round(W * 7.5 / 13.333)) // 2 * 2
    print('1/3 sprites: rendering every top-level shape alone (LibreOffice)…', flush=True)
    meta = build_sprite_decks(src, tmp, sel)
    blk = to_pngs(tmp / 'sprites_black.pptx', tmp / 'black', int(W * PAGE), profile)
    wht = to_pngs(tmp / 'sprites_white.pptx', tmp / 'white', int(W * PAGE), profile)
    base = to_pngs(tmp / 'base.pptx', tmp / 'base', W, profile)
    slides = []
    for i, si in enumerate(sel):
        recs = meta[si]
        sprites = []
        for r in recs:
            bi, wi = Image.open(blk[r['page']]), Image.open(wht[r['page']])
            sprites.append(Sprite(r, bi, wi, W, H))
        bimg = Image.open(base[i]).convert('RGBA').resize((W, H))
        root = etree.fromstring(z.read(parts[si - 1]))
        videos = {}
        for sp in sprites:
            if sp.video:
                mx, my = int(MARGIN / 13.333 * W), int(MARGIN / 7.5 * H)
                videos[sp.video] = VideoReader(z.read(sp.video), ((sp.img.width - 2 * mx) // 2 * 2, (sp.img.height - 2 * my) // 2 * 2), a.fps)
        slides.append(Slide(si, root, bimg, sprites, int(a.click_gap * 1000), videos))
    print(f'2/3 timing: {sum(len(s.effects) for s in slides)} effects on {len(slides)} slides', flush=True)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    ff = subprocess.Popen(['ffmpeg', '-v', 'error', '-y', '-f', 'rawvideo', '-pix_fmt', 'rgb24', '-s', f'{W}x{H}', '-r', str(a.fps), '-i', '-',
                           '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-crf', '20', '-preset', 'medium', '-movflags', '+faststart', str(out)],
                          stdin=subprocess.PIPE)
    try:
        font = ImageFont.truetype(subprocess.run(['fc-match', '-f', '%{file}', 'JetBrains Mono:bold'], capture_output=True, text=True).stdout,
                                  a.label_size or max(15, W // 60))
    except OSError:
        font = ImageFont.load_default()
    rng = np.random.default_rng(3)
    fps = a.fps
    total = 0
    log = []
    cap = int(a.max_seconds * fps) if a.max_seconds and a.max_seconds > 0 and a.cut else None
    holds = [natural_hold(S, a) for S in slides]
    trans = [S.trans_ms if i > 0 else 0 for i, S in enumerate(slides)]
    fitted, trans_k = fit_budget(holds, trans, a.max_seconds * 1000 if (a.max_seconds and not a.cut) else 0)
    if sum(fitted) < sum(holds) or trans_k < 1:
        print(f'    --max-seconds {a.max_seconds:g}: holds {sum(holds) / 1000:.0f} s -> {sum(fitted) / 1000:.0f} s, transitions x{trans_k:.2f}'
              ' so every slide is shown', flush=True)
    for i, S in enumerate(slides):
        if cap is not None and total >= cap:
            print(f'    --max-seconds {a.max_seconds:g} reached (--cut): stopped before slide {S.n}', flush=True)
            break
        hold_ms = fitted[i]
        S.trans_ms = int(S.trans_ms * trans_k)
        # transition into S from the previous slide
        if i > 0 and S.trans_ms > 0:
            P_ = slides[i - 1]
            ta = P_.hold_ms
            fa = P_.frame(ta, W, H)
            fb = S.frame(0, W, H)
            n = max(1, int(S.trans_ms / 1000 * fps))
            transition_frame._n = rng.random((H // 24 + 1, W // 24 + 1)).astype(np.float32)
            for f in range(n):
                if cap is not None and total >= cap:
                    break
                k = (f + 1) / (n + 1)
                img = morph_frame(P_, S, ta, k, W, H) if S.trans == 'morph' else transition_frame(P_, S, fa, fb, k, S.trans, S.trans_dir, W, H, rng)
                ff.stdin.write(tag_frame(img.convert('RGBA'), f'{a.label} · {S.trans}', font).convert('RGB').tobytes())
                total += 1
        S.hold_ms = hold_ms
        n = int(hold_ms / 1000 * fps)
        for f in range(n):
            if cap is not None and total >= cap:
                break
            img = S.frame(f * 1000 / fps, W, H)
            ff.stdin.write(tag_frame(img, a.label, font).convert('RGB').tobytes())
            total += 1
        log.append(f'slide {S.n:2}: {S.trans:<9} {S.trans_ms:5} ms, {len(S.effects):3} effects, hold {hold_ms / 1000:.1f} s')
        print('   ', log[-1], flush=True)
    ff.stdin.close()
    ff.wait()
    print(f'3/3 video: {out} ({total / fps:.1f} s, {W}x{H}, {fps} fps) — SIMULATION, label burnt in')
    (out.with_suffix('.log.txt')).write_text('\n'.join(log) + f'\nduration {total / fps:.1f} s\n', encoding='utf-8')
    if not a.keep:
        shutil.rmtree(tmp, ignore_errors=True)
    shutil.rmtree(profile, ignore_errors=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
