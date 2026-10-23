"""Build a 16:9 motion deck from a JSON spec: a designed system, named shapes ready for Morph.

  python3 build_deck.py deck.json --out deck/build.pptx [--base template.pptx] [--quiet]

Why: the look of a professional motion deck comes from a SYSTEM (ground, chrome, grid,
type roles, palette roles) applied identically on every slide, plus a few components
that move well (headings, cards, stats, screens, timelines, bars, lists, quotes). Write
the system once (a style pack), then write slides as compositions and keyframes.
Units are inches on a 13.333 x 7.5 canvas.

Top level
  "style": "path/to/style.json" (optional)   palette, fonts, type roles, grid, grounds, chrome, card...
  "palette" {key: "#RRGGBB"}  "fonts" {key: family}  "type" {role: {...}}  "grid" {...}  merged over the style
  "grounds": {"main": "assets/ground.jpg" | {recipe}}   "bg": default ground or colour for every slide
  "chrome": {"under": [shapes], "over": [shapes]}       replaces the style's chrome ("chrome_add" appends)
  "title": "Deck title"  (fills {title} in chrome)      "motion": {"default_ms": 1300, "stagger_ms": 220}
  "compat": "morph"|"safe"|"dual"   no-Morph safe mode (compat.py): safe rebuilds every Morph as classic
            transition + motion path/zoom/type effects; dual keeps Morph with a classic fallback
  "safe_transition": "fade"|{"type": "push", "dir": "l"}   transition that replaces Morph in safe mode
Slide keys
  "id", "from": "prev"|id, "set", "drop", "add", "front"   keyframes (see below)
  "bg": ground name | palette key | "#hex" | {"image": path}   "scene": "S3" (slide name, for audits)
  "bookend": true   a declared circular close: this slide may return to the opening image (check_deck allows
                    it on at most 2 slides instead of counting it as image reuse)
  "still": true     a deliberately still page (a printed chapter, a quiet opener): not counted in the moving-depth
                     share (at most a quarter of the deck)
  "chrome": true|false|[names]  "chrome_set": {name: {...}}  "source": "Nguồn: ..." (fills {source})
  "ms": 2000  "morph": "byObject"|"byWord"|"byChar"  "keep_timing": true   -> written to motion.json
  "notes": speaker notes
Slide keys also: "hidden": true drops the slide from the pptx but keeps it in deck.json (critic CUT).
Shapes (every shape: type, name, x, y, w, h | "col": [start, span] on the grid, rot, alpha, hidden, shadow)
  text     text (str | [paragraph | {text, size, color, ...}]), role (kicker, display, title, subtitle, body,
           caption, label, numeral, quote, cardtitle, pageno...), font, size, bold, italic, color, align l|c|r,
           anchor t|m|b, spacing (pt), line (multiple), caps, fade 0..1 (+ fade_bg: a quiet text as a solid
           colour mixed toward its background; prefer it to alpha, which LibreOffice renders badly)
  image    path, fit cover|contain|stretch, focus [fx, fy], radius, line_color, line_width,
           grayscale, duotone ["#dark", "#light"], brightness, contrast (-1..1)
  rect / ellipse   fill, alpha, gradient [[colour, alpha, pos%]...], angle, line_color, line_width,
           line_alpha, radius (0..0.5), geom (any MSO_SHAPE name: CHEVRON, TRIANGLE, ...)
  line     x1, y1, x2, y2, color, width, alpha, cap round|flat, dash dash|dot
  group    shapes [...]: moves and morphs as one object
Components (expand into named parts; set the component in a keyframe, Morph moves the parts)
  heading  kicker, title, sub, body, title_role, align           parts NAME.kick/.title/.sub/.body
  card     title, kicker, desc, image, icon, state rest|lit|dim    a group NAME with children "NAME title"...
  stat     value, unit, label, src, size, color, label_color        NAME.value/.unit/.label/.src
  screen   image, device browser|phone|plain, url, focus           a group NAME
  timeline nodes [{t, label, desc}], active                        NAME.rail/.prog/.dot{i}/.t{i}/.l{i}/.d{i}
  bars     data [{label, value, display}], max, grow 0..1, highlight   NAME.base/.bar{i}/.v{i}/.k{i}
  list     items [str | {text, sub}], reveal k, focus i, numbered  NAME.n{i}/.i{i}/.r{i}
  quote    text, by, mark_size (pt, default 2.2 x size)              NAME.mark/.text/.by
  callout  at [x, y], x, y, label, desc, color, chip, elbow, dot      NAME.ring/.dot/.lead/.labk/.lab/.desc (v5.4)
  swatch   image | colors, n, dir h|v, hex, labels, vivid            NAME.s{i}/.h{i}: a palette strip sampled from a photo
  photogrid images [path | {path, span [c, r], label, focus}], cols, rows, gap, border, tilt   NAME.f{i}/.p{i}/.t{i}k/.t{i}
  nodemap  nodes [{x, y (0..1), label, desc, side, big}], links [[a, b]], image, dash   NAME.map/.e{k}/.n{i}/.t{i}k/.t{i}/.d{i}
v5 motion and style keys
  slide  "style": "swiss-blue"   switch the whole system (palette, fonts, roles, grounds, chrome) for this
                                 slide and its keyframes: a deck can change style per chapter
         "transition": "curtains" | {"type": "push", "dir": "l", "ms": 800}   any native transition
                                 (motion.py --list); default Morph. "ms" still sets the duration.
  shape  "anim": {"fx": "fly", "dir": "l", "delay": 200, "ms": 600} | [ {...}, ... ]   entrance,
                                 emphasis, exit, path, orbit, letter-by-letter ("fx": "type"), loops
                                 ("loop": true), clicks ("click": 1) and triggers ("trigger": "<name>").
                                 On a component that expands into several top-level parts, the effect is
                                 applied to every part with "stagger" ms between them (default 120).
         "link": "<slide id>"    click the shape in the show to jump there (a hand-made slide zoom)
  line   "arrow": true | "stealth" puts an arrow head at x2,y2 ("arrow_start" at x1,y1)
  types  video  path (mp4), poster (png/jpg; default: first frame via ffmpeg), loop, mute, autoplay
         poly   points [[x, y], ...] (inches), closed (default true; "open": true = closed false), fill, alpha, line_color, line_width: freeform
                shapes for isometric solids, paper-cut layers, arrows and masks
Keyframes: "from": "prev" copies the slide; "set" {name: {keys}} changes shapes and components by
name; "drop" removes names; "add" appends; "front" moves names to the top. Unknown names are errors.
After building: use build.py (build_deck -> unpack -> motion -> pack -> render -> film_check).
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# v5.3.1 speed: python-pptx rescans the whole slide for a free shape id (an O(n^2) list search) and recomputes a
# group's extents after every child it adds. A 2,000-part Anatomy poster took ~30 s per slide. Ids are now issued
# from one scan per slide tree, and group extents are computed once, after the last child (see add_shape).
from pptx.oxml.shapes.groupshape import CT_GroupShape as _CTGroup

_NEXT_ID = {}


def _fast_next_shape_id(self):
    root = self.getroottree().getroot()
    key = id(root)
    held = _NEXT_ID.get(key)
    if held is None or held[0] is not root:
        used = [int(v) for v in root.xpath('//@id') if v.isdigit()]
        held = [root, max(used, default=0)]
        _NEXT_ID[key] = held
    held[1] += 1
    return held[1]


_CTGroup._next_shape_id = property(_fast_next_shape_id)
from pptx.shapes.shapetree import _BaseShapes as _PPTXShapes   # the second id source in python-pptx: share the counter
_PPTXShapes._next_shape_id = property(lambda self: _fast_next_shape_id(self._spTree))

HERE = Path(__file__).resolve().parent
STYLES = HERE.parent / 'assets' / 'styles'
W, H = 13.333, 7.5
DEFAULT_GRID = {'cols': 12, 'left': 0.9, 'right': 0.9, 'top': 0.8, 'bottom': 0.8, 'gutter': 0.25, 'safe': 0.4}
DEFAULT_TYPE = {
    'kicker': {'font': 'label', 'size': 12, 'bold': True, 'spacing': 3, 'caps': True, 'color': 'accent'},
    'display': {'font': 'display', 'size': 120, 'bold': True, 'line': 0.9, 'color': 'head'},
    'title': {'font': 'display', 'size': 44, 'bold': True, 'line': 0.95, 'color': 'head'},
    'subtitle': {'font': 'text', 'size': 24, 'line': 1.2, 'color': 'muted'},
    'body': {'font': 'text', 'size': 16, 'line': 1.3, 'color': 'ink'},
    'caption': {'font': 'text', 'size': 10, 'color': 'muted'},
    'label': {'font': 'label', 'size': 10, 'bold': True, 'spacing': 1.5, 'caps': True, 'color': 'muted'},
    'numeral': {'font': 'display', 'size': 260, 'bold': True, 'line': 0.85, 'color': 'accent'},
    'quote': {'font': 'display', 'size': 32, 'italic': True, 'line': 1.1, 'color': 'head'},
    'cardtitle': {'font': 'display', 'size': 24, 'bold': True, 'line': 1.0, 'color': 'head'},
    'pageno': {'font': 'text', 'size': 10, 'color': 'muted'},
}
# Neutral greyscale fallback only, so a deck that sets no palette looks unfinished rather than "styled".
# Every deck defines its own palette and fonts (principles/art-direction.md); nothing here is a house style.
DEFAULT_PALETTE = {'ground': '#FFFFFF', 'ground2': '#F2F2F2', 'surface': '#F7F7F7', 'ink': '#222222', 'head': '#111111',
                   'muted': '#777777', 'faint': '#DDDDDD', 'accent': '#444444', 'accent2': '#888888', 'lit': '#111111',
                   'veil': '#FFFFFF'}
DEFAULT_FONTS = {'display': 'DejaVu Sans', 'text': 'DejaVu Sans', 'label': 'DejaVu Sans'}
DEFAULT_CARD = {'fill': 'surface', 'fill_alpha': 1.0, 'line': 'faint', 'line_width': 0.75, 'radius': 0.05,
                'lit_fill': 'surface', 'lit_line': 'accent', 'rule': 'accent', 'dim_alpha': 0.66, 'shadow': None}
DEFAULT_SCREEN = {'bezel': '#16191E', 'line': '#2A2F37', 'line_width': 0.75, 'radius': 0.05, 'bar': '#1E2228'}
DEFAULT_SHADOW = {'blur': 20, 'dist': 6, 'dir': 90, 'color': '#000000', 'alpha': 0.35}
COMPONENTS = ('heading', 'card', 'stat', 'screen', 'timeline', 'bars', 'list', 'quote', 'callout', 'swatch', 'photogrid', 'nodemap',
              'construct')

def poly_closed(sh):
    """A poly is closed unless "closed": false or "open": true (both spellings are documented)."""
    if 'open' in sh:
        return not bool(sh['open'])
    return bool(sh.get('closed', True))

INHERIT = ('bg', 'scene', 'chrome', 'chrome_set', 'source', 'style')
STYLE_KEYS = ('palette', 'fonts', 'type', 'grid', 'grounds', 'chrome', 'chrome_add', 'card', 'screen', 'shadow', 'bg')


def load_style(name):
    """A style pack by path, or by name from assets/styles/ (empty by default: decks bring their own)."""
    p = Path(name)
    if not p.suffix:
        p = STYLES / f'{name}.json'
        if not p.is_file():
            p = STYLES.parent / 'styles-experimental' / f'{name}.json'
    if not p.is_file():
        known = sorted(x.stem for x in STYLES.glob('*.json'))
        raise SystemExit(f'style {name!r} not found. Known styles: {known} (or give a path to a style .json, or define it in '
                         f'deck.json: "styles": {{"{name}": {{"palette": {{...}}, "fonts": {{...}}}}}})')
    return json.loads(p.read_text(encoding='utf-8'))


def inline_style(name, data):
    """v5.4: a style defined inside deck.json. "styles": {"<name>": {...}} always defines one; a
    "style_overrides" entry defines one when <name> is not a shipped pack. Either may say
    "extends": "<pack or path>" and give only what differs; without "extends" it is a full spec."""
    if not name or Path(str(name)).suffix:
        return None
    spec = (data.get('styles') or {}).get(name)
    if spec is None and not (STYLES / f'{name}.json').is_file() and not (STYLES.parent / 'styles-experimental' / f'{name}.json').is_file():
        spec = (data.get('style_overrides') or {}).get(name)
    if spec is None:
        return None
    spec = copy.deepcopy(spec)
    base = spec.pop('extends', None)
    return merge(load_style(base), spec) if base else spec


def merge(a, b):
    out = copy.deepcopy(a)
    for k, v in (b or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = merge(out[k], v)
        else:
            out[k] = copy.deepcopy(v)
    return out


class Spec:
    def __init__(self, data, root):
        style_ref = data.get('style')
        if style_ref and Path(str(style_ref)).suffix and not Path(str(style_ref)).is_absolute():
            style_ref = str(root / style_ref)
        st = inline_style(style_ref, data)
        if st is None:
            st = load_style(style_ref) if style_ref else {}
        self.style_name = Path(str(data.get('style', 'custom'))).stem
        self.style = st
        self.pal = {**DEFAULT_PALETTE, **st.get('palette', {}), **data.get('palette', {})}
        self.fonts = {**DEFAULT_FONTS, **st.get('fonts', {}), **data.get('fonts', {})}
        self.type = merge(merge(DEFAULT_TYPE, st.get('type', {})), data.get('type', {}))
        self.grid = {**DEFAULT_GRID, **st.get('grid', {}), **data.get('grid', {})}
        self.grounds = {**st.get('grounds', {}), **data.get('grounds', {})}
        chrome = data['chrome'] if 'chrome' in data else st.get('chrome', {})
        chrome = copy.deepcopy(chrome or {})
        for layer, shapes in (data.get('chrome_add') or {}).items():
            chrome.setdefault(layer, []).extend(copy.deepcopy(shapes))
        self.chrome = chrome
        self.card = {**DEFAULT_CARD, **st.get('card', {}), **data.get('card', {})}
        self.screen = {**DEFAULT_SCREEN, **st.get('screen', {}), **data.get('screen', {})}
        self.shadow = {**DEFAULT_SHADOW, **(st.get('shadow') or {}), **(data.get('shadow') or {})}
        self.motion = {**{'default_ms': 1300, 'stagger_ms': 220}, **st.get('motion', {}), **data.get('motion', {})}
        self.default_bg = data.get('bg', st.get('default_bg'))
        self.title = data.get('title', '')
        self.raw = data
        self.root = root
        self.warnings = []
        self._ground_cache = {}
        self._data = data
        self._subs = {}

    def for_style(self, name):
        """The system for one slide: the deck's own style, or another pack for a chapter.
        Deck-level palette/fonts/type overrides belong to the deck style; a chapter style takes its
        own overrides from "style_overrides": {"<style>": {"palette": {...}, ...}}."""
        if not name or name == self._data.get('style'):
            return self
        if name not in self._subs:
            d = {k: v for k, v in self._data.items() if k not in STYLE_KEYS}
            d['style'] = name
            if inline_style(name, self._data) is None:   # an inline style already carries its own values
                d.update(copy.deepcopy((self._data.get('style_overrides') or {}).get(name, {})))
            sub = Spec(d, self.root)
            sub.warnings = self.warnings
            self._subs[name] = sub
        return self._subs[name]

    # -- values
    def color(self, c):
        c = self.pal.get(c, c)
        if not (isinstance(c, str) and re.fullmatch(r'#[0-9A-Fa-f]{6}', c)):
            raise SystemExit(f'bad colour {c!r}: use "#RRGGBB" or a palette key {sorted(self.pal)}')
        return c[1:].upper()

    def font(self, f):
        return self.fonts.get(f, f)

    def path(self, p):
        q = Path(p)
        return q if q.is_absolute() else (self.root / q)

    def role(self, name):
        if not name:
            return {}
        # v5.6 (I29): component-scoped roles: "timeline.cardtitle" inherits "cardtitle" and overrides only what it sets,
        # so enlarging card titles (deck "type": {"cardtitle": ...}) need not enlarge the timeline's labels
        if '.' in name:
            base = name.split('.', 1)[1]
            if base not in self.type:
                raise SystemExit(f'unknown text role {name!r}: {base!r} is not a role; known: {sorted(self.type)}')
            return {**self.type[base], **self.type.get(name, {})}
        if name not in self.type:
            raise SystemExit(f'unknown text role {name!r}; known: {sorted(self.type)}')
        return self.type[name]

    def col(self, start, span):
        g = self.grid
        colw = (W - g['left'] - g['right'] - (g['cols'] - 1) * g['gutter']) / g['cols']
        return g['left'] + (start - 1) * (colw + g['gutter']), span * colw + (span - 1) * g['gutter']

    def ground(self, name):
        """Image path for a ground name, generating a style recipe on first use; None if not a ground."""
        if name not in self.grounds:
            return None
        g = self.grounds[name]
        if isinstance(g, str):
            p = self.path(g)
            if not p.exists():
                raise SystemExit(f'ground {name!r}: {p} not found')
            return p
        if name in self._ground_cache:
            return self._ground_cache[name]
        import surface                                  # recipe grounds: scripts/surface.py (principles/surface-and-depth.md)
        key = f'{self.style_name}-{name}'
        out = self.root / 'assets' / 'grounds' / f'{key}.jpg'
        if not out.exists():
            out.parent.mkdir(parents=True, exist_ok=True)
            rec = {k: (tuple(v) if isinstance(v, list) else v) for k, v in g.items()}
            base = rec.pop('base')
            base = base if str(base).startswith('#') else self.pal[base]
            rec['light'] = [tuple(x) for x in rec.get('light', [])]
            surface.ground(rec.pop('w', 2400), rec.pop('h', 1350), base, rec.pop('seed', 1), **rec).save(out, quality=88)
        self._ground_cache[name] = out
        return out


# ---------------------------------------------------------------- keyframes

def resolve(slides):
    """Expand from/set/drop/add into full shape lists."""
    out, by_id = [], {}
    for i, s in enumerate(slides):
        s = copy.deepcopy(s)
        src = s.get('from')
        if src:
            base = out[-1] if src == 'prev' else by_id.get(src)
            if base is None:
                raise SystemExit(f'slide {i + 1}: "from": {src!r} not found')
            shapes = copy.deepcopy(base['shapes'])
            for k in INHERIT:
                if k in base and k not in s:
                    s[k] = copy.deepcopy(base[k])
            names = [sh.get('name') for sh in shapes]
            for n, over in s.get('set', {}).items():
                if n not in names:
                    hint = ''
                    if '.' in n and n.split('.')[0] in names:
                        hint = f' (it is a part of component {n.split(".")[0]!r}: set the component\'s keys instead)'
                    raise SystemExit(f'slide {i + 1}: set {n!r}: no shape with that name in the source slide{hint}')
                shapes[names.index(n)].update(over)
            drop = set(s.get('drop', []))
            missing = drop - set(names)
            if missing:
                raise SystemExit(f'slide {i + 1}: drop {sorted(missing)}: not in the source slide')
            shapes = [sh for sh in shapes if sh.get('name') not in drop]
            shapes += copy.deepcopy(s.get('add', []))
            s['shapes'] = shapes
        for n in s.get('front', []):
            idx = [k for k, sh in enumerate(s['shapes']) if sh.get('name') == n]
            for k in idx:
                s['shapes'].append(s['shapes'].pop(k))
        s.setdefault('shapes', [])
        names = [sh.get('name') for sh in s['shapes'] if (sh.get('name') or '').startswith('!!')]
        dup = {n for n in names if names.count(n) > 1}
        if dup:
            raise SystemExit(f'slide {i + 1}: duplicate !! names {sorted(dup)} (Morph pairs by name: must be unique per slide)')
        out.append(s)
        if s.get('id'):
            by_id[s['id']] = s
    return out


# ---------------------------------------------------------------- text sizing

def text_height(spec, sh, w):
    """Estimated height (in) of a text shape laid out in width w."""
    import typeset
    paras = sh['text'] if isinstance(sh['text'], list) else [sh['text']]
    total = 0.0
    for para in paras:
        d = para if isinstance(para, dict) else {'text': para}
        p = text_props(spec, {**sh, **{k: v for k, v in d.items() if k != 'text'}})
        txt = d.get('text') if 'runs' not in d else ''.join(r.get('text', '') for r in d['runs'])
        m = typeset.measure(txt or ' ', p['font'], p['size'], w, p['bold'], p['italic'], p['line'], p['spacing'], p['caps'])
        if m['font_missing']:
            chars = max(1, int(w * 72 / (p['size'] * 0.52)))
            lines = sum(max(1, -(-len(line) // chars)) for line in (txt or ' ').split('\n'))
            total += lines * p['size'] * 1.2 * p['line'] / 72
        else:
            total += m['height_in']
        total += d.get('space_before', 0) / 72
    return total


def text_props(spec, sh):
    r = spec.role(sh.get('role'))
    g = lambda k, d=None: sh.get(k, r.get(k, d))  # noqa: E731
    return {
        'font': spec.font(g('font', 'text')), 'size': g('size', 18), 'bold': bool(g('bold', False)),
        'italic': bool(g('italic', False)), 'color': g('color', 'ink'), 'spacing': g('spacing', 0) or 0,
        'line': g('line', 1.0) or 1.0, 'caps': bool(g('caps', False)), 'alpha': g('alpha'),
        'align': g('align', 'l'), 'anchor': g('anchor', 't'), 'fade': g('fade'), 'fade_bg': g('fade_bg', 'ground'),
    }


def mix(spec, fg, bg, amount):
    """fg at `amount` opacity over bg, as one solid colour (hex without #)."""
    a, b = spec.color(fg), spec.color(bg)
    out = ''.join(f'{round(int(a[i:i + 2], 16) * amount + int(b[i:i + 2], 16) * (1 - amount)):02X}' for i in (0, 2, 4))
    return '#' + out


# ---------------------------------------------------------------- components

def _nm(base, part):
    return f'{base}.{part}' if base else None


def _child(base, part):
    return f'{base} {part}' if base else None


def comp_heading(spec, sh):
    x, y, w = sh['x'], sh['y'], sh['w']
    align = sh.get('align', 'l')
    parts = [('kick', 'kicker', sh.get('kicker'), 0.14), ('title', sh.get('title_role', 'title'), sh.get('title'), 0.22),
             ('sub', 'subtitle', sh.get('sub'), 0.26), ('body', 'body', sh.get('body'), 0.0)]
    out, cy = [], y
    for key, role, txt, gap in parts:
        if not txt:
            continue
        t = {'type': 'text', 'name': _nm(sh.get('name'), key), 'role': role, 'text': txt, 'x': x, 'y': cy, 'w': w, 'align': align}
        t.update(sh.get(f'{key}_style', {}))
        if sh.get('alpha') is not None:
            t['alpha'] = sh['alpha']
        h = text_height(spec, t, w) + 0.04
        t['h'] = round(h, 3)
        out.append(t)
        cy += h + gap
    return out


def comp_card(spec, sh):
    c = spec.card
    x, y, w, h = sh['x'], sh['y'], sh['w'], sh['h']
    state = sh.get('state', 'rest')
    base = sh.get('name')
    pad = sh.get('pad', 0.24)
    lit = state == 'lit'
    dim = state == 'dim'
    surf = {'type': 'rect', 'name': _child(base, 'bg'), 'x': x, 'y': y, 'w': w, 'h': h, 'radius': c['radius'],
            'fill': c['lit_fill'] if lit else c['fill'], 'alpha': 1.0 if lit else c['fill_alpha'],
            'line_color': c['lit_line'] if lit else c['line'], 'line_width': c['line_width'] * (1.4 if lit else 1)}
    if not surf['line_color']:  # an open card: no outline at rest, only the rule (the reference's role cards)
        surf.pop('line_color')
    if not surf['alpha']:
        surf['alpha'] = 0.001  # keep a real fill so Morph can fade the surface in when the card lights up
    shadow = c.get('lit_shadow') if lit and c.get('lit_shadow') else c.get('shadow')
    if shadow:
        surf['shadow'] = shadow
    kids = [surf]
    cy = y + pad
    ta = c['dim_alpha'] if dim else None
    fade_bg = c['fill'] if (c.get('fill_alpha', 1) or 0) >= 0.5 else sh.get('ground', 'ground')
    if sh.get('image'):
        ih = sh.get('image_h', h * 0.42)
        kids.append({'type': 'image', 'name': _child(base, 'image'), 'path': sh['image'], 'x': x + pad, 'y': cy,
                     'w': w - 2 * pad, 'h': ih, 'fit': 'cover', 'focus': sh.get('focus', [0.5, 0.5]),
                     'radius': 0.04, 'alpha': 0.5 if dim else None})
        cy += ih + 0.2
    elif sh.get('icon'):
        ih = sh.get('icon_h', min(1.3, h * 0.3))
        kids.append({'type': 'image', 'name': _child(base, 'icon'), 'path': sh['icon'], 'x': x + pad, 'y': cy,
                     'w': w - 2 * pad, 'h': ih, 'fit': 'contain', 'alpha': 0.5 if dim else None})
        cy += ih + 0.2
    kids.append({'type': 'line', 'name': _child(base, 'rule'), 'x1': x + pad, 'y1': cy, 'x2': x + pad + min(0.9, w * 0.3), 'y2': cy,
                 'color': c['rule'] if not dim else (c['line'] or 'faint'), 'width': 1.5 if lit else 1.0})
    cy += 0.14
    for part, role, key in (('title', 'cardtitle', 'title'), ('kicker', 'label', 'kicker'), ('desc', 'body', 'desc')):
        if not sh.get(key):
            continue
        t = {'type': 'text', 'name': _child(base, part), 'role': role, 'text': sh[key], 'x': x + pad, 'y': cy, 'w': w - 2 * pad}
        if part == 'desc':
            t['size'] = sh.get('desc_size', max(11, spec.role('body').get('size', 15) - 2))
        if part == 'kicker':
            t['color'] = spec.role('kicker').get('color', c['rule']) if not dim else 'muted'
        if ta is not None and part != 'kicker':  # the small label stays legible: it is already muted
            t['fade'], t['fade_bg'] = ta, fade_bg
        t['h'] = round(text_height(spec, t, w - 2 * pad) + 0.04, 3)
        kids.append(t)
        cy += t['h'] + (0.1 if part != 'desc' else 0)
    if cy > y + h + 0.05:
        spec.warnings.append(f'{base or "card"}: content needs {cy - y:.2f} in but the card is {h:.2f} in tall')
    return [{'type': 'group', 'name': base, 'shapes': kids}]


def comp_stat(spec, sh):
    x, y, w = sh['x'], sh['y'], sh['w']
    base = sh.get('name')
    num = {'type': 'text', 'name': _nm(base, 'value'), 'role': sh.get('role', 'numeral'), 'text': str(sh['value']),
           'x': x, 'y': y, 'w': w}
    if sh.get('size'):
        num['size'] = sh['size']
    if sh.get('color'):
        num['color'] = sh['color']
    p = text_props(spec, num)
    import typeset
    m = typeset.measure(str(sh['value']), p['font'], p['size'], w, p['bold'], p['italic'], p['line'], p['spacing'], p['caps'])
    vw = m['widest_in'] if not m['font_missing'] else len(str(sh['value'])) * p['size'] * 0.55 / 72
    vh = text_height(spec, num, w)
    num['h'] = round(vh + 0.02, 3)
    out = [num]
    if sh.get('unit'):
        avail = max(0.6, w - vw - 0.12)
        us = max(16, int(p['size'] * 0.3))
        probe = {'type': 'text', 'role': 'title', 'text': sh['unit']}
        while us > 16:
            tp = text_props(spec, {**probe, 'size': us})
            um = typeset.measure(sh['unit'], tp['font'], us, 99, tp['bold'], tp['italic'], tp['line'], tp['spacing'], tp['caps'])
            if um['font_missing'] or um['widest_in'] <= avail:
                break
            us -= 2
        out.append({'type': 'text', 'name': _nm(base, 'unit'), 'role': 'title', 'size': us, 'text': sh['unit'],
                    'x': round(x + vw + 0.12, 3), 'y': round(y + vh * 0.62 - us / 72 * 1.1, 3), 'w': round(avail, 3),
                    'h': round(us / 72 * 1.35, 3), 'color': sh.get('unit_color', p['color'])})
    cy = y + vh + 0.12
    if sh.get('label'):
        t = {'type': 'text', 'name': _nm(base, 'label'), 'role': 'body', 'text': sh['label'], 'x': x, 'y': cy, 'w': w}
        if sh.get('label_color'):
            t['color'] = sh['label_color']
        t['h'] = round(text_height(spec, t, w) + 0.04, 3)
        out.append(t)
        cy += t['h'] + 0.06
    if sh.get('src'):
        t = {'type': 'text', 'name': _nm(base, 'src'), 'role': 'caption', 'text': sh['src'], 'x': x, 'y': cy, 'w': w}
        if sh.get('label_color'):
            t['color'] = sh['label_color']
        t['h'] = round(text_height(spec, t, w) + 0.03, 3)
        out.append(t)
    return out


def comp_screen(spec, sh):
    s = spec.screen
    x, y, w, h = sh['x'], sh['y'], sh['w'], sh['h']
    base = sh.get('name')
    device = sh.get('device', 'browser')
    kids = []
    if device == 'phone':
        r = 0.16
        kids.append({'type': 'rect', 'name': _child(base, 'bezel'), 'x': x, 'y': y, 'w': w, 'h': h, 'fill': s['bezel'], 'radius': r,
                     'line_color': s['line'], 'line_width': s['line_width'], 'shadow': sh.get('shadow', spec.shadow)})
        pad = w * 0.045
        ix, iy, iw, ih = x + pad, y + pad, w - 2 * pad, h - 2 * pad
        kids.append({'type': 'image', 'name': _child(base, 'image'), 'path': sh['image'], 'x': ix, 'y': iy, 'w': iw, 'h': ih,
                     'fit': 'cover', 'focus': sh.get('focus', [0.5, 0.0]), 'radius': 0.12})
        kids.append({'type': 'rect', 'name': _child(base, 'notch'), 'x': x + w * 0.36, 'y': y + pad + 0.06, 'w': w * 0.28,
                     'h': max(0.08, w * 0.045), 'fill': '#000000', 'radius': 0.5})
    elif device == 'plain':
        kids.append({'type': 'image', 'name': _child(base, 'image'), 'path': sh['image'], 'x': x, 'y': y, 'w': w, 'h': h,
                     'fit': 'cover', 'focus': sh.get('focus', [0.5, 0.5]), 'radius': s['radius'],
                     'line_color': s['line'], 'line_width': s['line_width'], 'shadow': sh.get('shadow', spec.shadow)})
    else:
        bar = min(0.34, h * 0.09)
        kids.append({'type': 'rect', 'name': _child(base, 'bezel'), 'x': x, 'y': y, 'w': w, 'h': h, 'fill': s['bezel'],
                     'radius': s['radius'] * 0.6, 'line_color': s['line'], 'line_width': s['line_width'],
                     'shadow': sh.get('shadow', spec.shadow)})
        kids.append({'type': 'rect', 'name': _child(base, 'bar'), 'x': x, 'y': y, 'w': w, 'h': bar,
                     'fill': s.get('bar', s['bezel']), 'radius': 0.0})
        for i in range(3):
            d = bar * 0.32
            kids.append({'type': 'ellipse', 'name': _child(base, f'dot{i}'), 'x': x + 0.16 + i * (d + 0.07), 'y': y + (bar - d) / 2,
                         'w': d, 'h': d, 'fill': s['line']})
        if sh.get('url'):
            kids.append({'type': 'text', 'name': _child(base, 'url'), 'role': 'caption', 'text': sh['url'],
                         'x': x + w * 0.3, 'y': y + (bar - 0.2) / 2, 'w': w * 0.4, 'h': 0.2, 'align': 'c', 'anchor': 'm'})
        kids.append({'type': 'image', 'name': _child(base, 'image'), 'path': sh['image'], 'x': x, 'y': y + bar, 'w': w,
                     'h': h - bar, 'fit': 'cover', 'focus': sh.get('focus', [0.5, 0.0])})
    return [{'type': 'group', 'name': base, 'shapes': kids}]


def scoped(spec, sh, comp, role):
    """v5.6 (I29): a component's role: the shape's "roles" map, else "<comp>.<role>" when the deck/style defines it."""
    r = (sh.get('roles') or {}).get(role)
    if r:
        return r
    return f'{comp}.{role}' if f'{comp}.{role}' in spec.type else role


def comp_timeline(spec, sh):
    x, y, w = sh['x'], sh['y'], sh['w']
    base = sh.get('name')
    nodes = sh['nodes']
    act = sh.get('active', 0)
    n = len(nodes)
    pos = [x + (w * (i / (n - 1)) if n > 1 else w / 2) for i in range(n)]
    lab_w = sh.get('label_w', min(2.3, w / max(1, n) - 0.15))
    color = sh.get('color', 'accent')
    out = [{'type': 'line', 'name': _nm(base, 'rail'), 'x1': x - 0.25, 'y1': y, 'x2': x + w + 0.25, 'y2': y,
            'color': sh.get('rail_color', 'faint'), 'width': 1.25}]
    end = pos[act] if act is not None and act >= 0 else x
    out.append({'type': 'line', 'name': _nm(base, 'prog'), 'x1': x - 0.25, 'y1': y, 'x2': max(x - 0.2, end), 'y2': y,
                'color': color, 'width': 2.25})
    for i, (px, nd) in enumerate(zip(pos, nodes)):
        on = i == act
        done = act is not None and i < act
        d = 0.42 if on else 0.28
        al = 'l' if (i == 0 or n == 1) else ('r' if i == n - 1 else 'c')
        lx = round(px - 0.05 if al == 'l' else (px + 0.05 - lab_w if al == 'r' else px - lab_w / 2), 3)
        out.append({'type': 'ellipse', 'name': _nm(base, f'dot{i}'), 'x': round(px - d / 2, 3), 'y': round(y - d / 2, 3), 'w': d, 'h': d,
                    'fill': color if (on or done) else 'ground', 'line_color': color, 'line_width': 1.25})
        ty = y + 0.38
        if nd.get('t'):
            t = {'type': 'text', 'name': _nm(base, f't{i}'), 'role': scoped(spec, sh, 'timeline', 'cardtitle'), 'text': nd['t'], 'align': al,
                 'x': lx, 'y': round(ty, 3), 'w': lab_w, 'color': 'head' if on else 'muted'}
            t['h'] = round(text_height(spec, t, lab_w) + 0.04, 3)
            out.append(t)
            ty += t['h'] + 0.04
        if nd.get('label'):
            t = {'type': 'text', 'name': _nm(base, f'l{i}'), 'role': scoped(spec, sh, 'timeline', 'body'), 'bold': True, 'text': nd['label'], 'align': al,
                 'x': lx, 'y': round(ty, 3), 'w': lab_w, 'color': 'ink' if on else 'muted'}
            t['h'] = round(text_height(spec, t, lab_w) + 0.04, 3)
            out.append(t)
            ty += t['h'] + 0.04
        if nd.get('desc'):
            t = {'type': 'text', 'name': _nm(base, f'd{i}'), 'role': scoped(spec, sh, 'timeline', 'body'), 'size': max(11, spec.role(scoped(spec, sh, 'timeline', 'body')).get('size', 15) - 2),
                 'text': nd['desc'], 'align': al, 'x': lx, 'y': round(ty, 3), 'w': lab_w, 'color': 'ink' if on else 'muted'}
            t['h'] = round(text_height(spec, t, lab_w) + 0.04, 3)
            out.append(t)
    return out


def comp_bars(spec, sh):
    x, y, w, h = sh['x'], sh['y'], sh['w'], sh['h']
    base = sh.get('name')
    data = sh['data']
    n = len(data)
    gap = sh.get('gap', 0.45)
    bw = w / (n + (n - 1) * gap)
    vmax = sh.get('max') or max(d['value'] for d in data) or 1
    grow = max(0.0, min(1.0, sh.get('grow', 1.0)))
    hl = sh.get('highlight')
    label_h = 0.36
    value_h = 0.5
    base_y = y + h - label_h
    plot_h = h - label_h - value_h
    out = [{'type': 'line', 'name': _nm(base, 'base'), 'x1': x - 0.1, 'y1': base_y, 'x2': x + w + 0.1, 'y2': base_y,
            'color': sh.get('axis_color', 'faint'), 'width': 1.0}]
    for i, d in enumerate(data):
        bx = x + i * bw * (1 + gap)
        bh = max(0.01, plot_h * d['value'] / vmax * grow)
        col = sh.get('color', 'accent') if (hl is None or hl == i) else sh.get('muted_color', 'faint')
        out.append({'type': 'rect', 'name': _nm(base, f'bar{i}'), 'x': round(bx, 3), 'y': round(base_y - bh, 3), 'w': round(bw, 3),
                    'h': round(bh, 3), 'fill': col})
        if grow >= 0.999:
            v = {'type': 'text', 'name': _nm(base, f'v{i}'), 'role': scoped(spec, sh, 'bars', 'cardtitle'), 'size': sh.get('value_size', spec.role(scoped(spec, sh, 'bars', 'cardtitle')).get('size', 22)),
                 'text': d.get('display', str(d['value'])), 'x': round(bx - 0.3, 3), 'w': round(bw + 0.6, 3), 'align': 'c', 'anchor': 'b',
                 'color': 'head' if (hl is None or hl == i) else 'muted'}
            vh = round(text_height(spec, v, v['w']) + 0.04, 3)
            v.update({'y': round(base_y - bh - vh - 0.06, 3), 'h': vh})
            out.append(v)
        out.append({'type': 'text', 'name': _nm(base, f'k{i}'), 'role': 'label', 'text': d['label'], 'x': round(bx - 0.3, 3),
                    'y': round(base_y + 0.1, 3), 'w': round(bw + 0.6, 3), 'h': 0.24, 'align': 'c'})
    return out


def comp_list(spec, sh):
    x, y, w = sh['x'], sh['y'], sh['w']
    base = sh.get('name')
    items = sh['items']
    k = sh.get('reveal', len(items))
    focus = sh.get('focus')
    numbered = sh.get('numbered', True)
    gap = sh.get('gap', 0.22)
    size = sh.get('size', max(16, spec.role('body').get('size', 16) + 2))
    out, cy = [], y
    num_w = 0.62 if numbered else 0
    for i, it in enumerate(items[:k]):
        d = it if isinstance(it, dict) else {'text': it}
        quiet = focus is not None and focus != i
        if numbered:
            out.append({'type': 'text', 'name': _nm(base, f'n{i}'), 'role': 'kicker', 'text': f'{i + 1:02d}', 'x': x, 'y': round(cy + 0.06, 3),
                        'w': num_w, 'h': 0.3, **({'color': 'muted'} if quiet else {})})
        paras = [{'text': d['text'], 'size': size, 'bold': True}]
        if d.get('sub'):
            paras.append({'text': d['sub'], 'size': max(11, size - 4), 'color': 'muted', 'space_before': 3, 'bold': False})
        t = {'type': 'text', 'name': _nm(base, f'i{i}'), 'role': 'body', 'text': paras, 'x': x + num_w, 'y': round(cy, 3),
             'w': w - num_w, 'fade': 0.66 if quiet else None, 'fade_bg': sh.get('ground', 'ground')}
        t['h'] = round(text_height(spec, t, w - num_w) + 0.04, 3)
        out.append(t)
        cy += t['h'] + gap / 2
        if sh.get('rules', True):
            out.append({'type': 'line', 'name': _nm(base, f'r{i}'), 'x1': x, 'y1': round(cy, 3), 'x2': x + w, 'y2': round(cy, 3),
                        'color': 'faint', 'width': 0.75})
        cy += gap / 2
    return out


def comp_quote(spec, sh):
    x, y, w = sh['x'], sh['y'], sh['w']
    base = sh.get('name')
    q = spec.role('quote')
    size = sh.get('size', q.get('size', 32))
    ms = int(sh.get('mark_size', size * 2.2))   # mark_size: put the mark on the deck's type scale
    mw = ms * 0.5 / 72
    out = [{'type': 'text', 'name': _nm(base, 'mark'), 'role': 'quote', 'text': '“', 'size': ms, 'italic': False,
            'x': round(x - mw - 0.08, 3), 'y': round(y - ms * 0.18 / 72, 3), 'w': round(mw, 3), 'h': round(ms * 1.25 / 72, 3),
            'align': 'r', 'color': sh.get('mark_color', 'accent'), 'bleed': True}]
    t = {'type': 'text', 'name': _nm(base, 'text'), 'role': 'quote', 'text': sh['text'], 'x': x, 'y': y, 'w': w}
    if sh.get('size'):
        t['size'] = sh['size']
    t['h'] = round(text_height(spec, t, w) + 0.06, 3)
    out.append(t)
    if sh.get('by'):
        out.append({'type': 'text', 'name': _nm(base, 'by'), 'role': 'label', 'text': '— ' + sh['by'], 'x': x, 'y': round(y + t['h'] + 0.18, 3),
                    'w': w, 'h': 0.3})
    return out


# ---------------------------------------------------------------- v5.4 magazine / collage components

def _role_or(spec, name, fallback):
    return name if name in spec.type else fallback


def _chip_text(spec, base, part, text, x, y, role, color, fill, pad=0.07, w=None, align='l', alpha=None):
    """A label on its own colour band: [rect, text]. The band is sized to the measured text."""
    t = {'type': 'text', 'name': _nm(base, part), 'role': role, 'text': text, 'x': x, 'y': y, 'w': w or 4.0, 'h': 0.5,
         'color': color, 'align': align}
    wide, need = text_extent(spec, t)
    t['w'] = round(w or (wide + 0.04), 3)
    t['h'] = round(need + 0.02, 3)
    band = {'type': 'rect', 'name': _nm(base, part + 'k'), 'x': round(x - pad, 3), 'y': round(y - pad * 0.8, 3),
            'w': round(t['w'] + 2 * pad, 3), 'h': round(t['h'] + 1.6 * pad, 3), 'fill': fill}
    if alpha is not None:
        band['alpha'] = alpha
    return [band, t]


def comp_callout(spec, sh):
    """A leader-line callout: a dot ON the object part, a hairline leader, and the label on a colour chip.
    at [x, y] = the point on the object (inches); x, y = where the label goes. NAME.ring/.dot/.lead/.labk/.lab/.desc"""
    base = sh.get('name')
    ax, ay = sh['at']
    col = sh.get('color', 'accent')
    r = float(sh.get('dot', 0.09))
    role = _role_or(spec, sh.get('role', 'callout'), 'label')
    chip = sh.get('chip', True)
    parts = _chip_text(spec, base, 'lab', sh['label'], sh['x'], sh['y'], role,
                       sh.get('text_color', 'surface' if chip else col), sh.get('chip_fill', col), w=sh.get('w'))
    if not chip:
        parts = parts[1:]
    band = parts[0]
    if sh.get('desc'):
        d = {'type': 'text', 'name': _nm(base, 'desc'), 'role': _role_or(spec, 'caption', 'body'), 'text': sh['desc'],
             'x': sh['x'], 'y': round(band['y'] + band['h'] + 0.06, 3), 'w': sh.get('desc_w', max(1.6, band['w'])), 'h': 0.4,
             'color': sh.get('desc_color', 'ink')}
        d['h'] = round(text_height(spec, d, d['w']) + 0.04, 3)
        parts.append(d)
    # the leader ends on the nearest edge of the chip (left or right, at mid height)
    bx, by, bw, bh = band['x'], band['y'], band['w'], band['h']
    ex = bx if ax <= bx else (bx + bw if ax >= bx + bw else ax)
    ey = by + bh / 2 if ex != ax else (by if ay < by else by + bh)
    lead = []
    if sh.get('elbow') and ex != ax:
        lead = [{'type': 'line', 'name': _nm(base, 'lead'), 'x1': ax, 'y1': ay, 'x2': ax, 'y2': round(ey, 3), 'color': col,
                 'width': sh.get('width', 1.25)},
                {'type': 'line', 'name': _nm(base, 'lead2'), 'x1': ax, 'y1': round(ey, 3), 'x2': round(ex, 3), 'y2': round(ey, 3),
                 'color': col, 'width': sh.get('width', 1.25)}]
    else:
        lead = [{'type': 'line', 'name': _nm(base, 'lead'), 'x1': ax, 'y1': ay, 'x2': round(ex, 3), 'y2': round(ey, 3), 'color': col,
                 'width': sh.get('width', 1.25)}]
    ring = {'type': 'ellipse', 'name': _nm(base, 'ring'), 'x': round(ax - 2.2 * r, 3), 'y': round(ay - 2.2 * r, 3),
            'w': round(4.4 * r, 3), 'h': round(4.4 * r, 3), 'fill': col, 'alpha': 0.22}
    dot = {'type': 'ellipse', 'name': _nm(base, 'dot'), 'x': round(ax - r, 3), 'y': round(ay - r, 3), 'w': round(2 * r, 3),
           'h': round(2 * r, 3), 'fill': col, 'line_color': sh.get('dot_line', 'surface'), 'line_width': 1.5}
    return lead + [ring, dot] + parts


SWATCH_CACHE = {}


def sample_palette(path, n=5, min_dist=42, vivid=False):
    """n dominant, mutually distinct colours of a picture (median-cut, then greedy by frequency)."""
    key = (str(path), n, min_dist, vivid)
    if key in SWATCH_CACHE:
        return SWATCH_CACHE[key]
    from PIL import ImageOps
    im = Image.open(path)
    try:
        im.seek(0)
    except EOFError:
        pass
    im = ImageOps.exif_transpose(im)
    if 'A' in im.getbands() or im.mode == 'P':
        # a cut-out: sample the object only, never the transparent page around it
        rgba = im.convert('RGBA')
        rgba.thumbnail((240, 240))
        import numpy as _np
        arr = _np.asarray(rgba)
        px = arr[arr[..., 3] >= 160][:, :3]
        im = Image.fromarray(px.reshape(1, -1, 3).astype('uint8'), 'RGB') if len(px) else Image.new('RGB', (1, 1))
    else:
        im = im.convert('RGB')
        im.thumbnail((200, 200))
    q = im.quantize(colors=48 if vivid else max(12, n * 4), method=Image.Quantize.MEDIANCUT)
    pal = q.getpalette()
    counts = sorted(q.getcolors(), reverse=True)
    if vivid:   # weight saturated colours up: the red trim or the gold foil, not five greys
        sat = lambda i: (max(pal[i * 3:i * 3 + 3]) - min(pal[i * 3:i * 3 + 3])) / 255  # noqa: E731
        counts = sorted(counts, key=lambda ci: -ci[0] * (0.15 + sat(ci[1]) ** 1.5 * 4))
    picked = []
    for _, idx in counts:
        c = tuple(pal[idx * 3:idx * 3 + 3])
        if all(sum((a - b) ** 2 for a, b in zip(c, p)) ** 0.5 >= min_dist for p in picked):
            picked.append(c)
        if len(picked) == n:
            break
    out = ['#%02X%02X%02X' % c for c in picked]
    SWATCH_CACHE[key] = out
    return out


def _ink_on(hexc):
    r, g, b = (int(hexc[i:i + 2], 16) / 255 for i in (1, 3, 5))
    lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
    return '#1A1410' if lum > 0.5 else '#FFFFFF'


def comp_swatch(spec, sh):
    """A palette strip sampled from a photo (or given colours), with hex codes under each chip.
    image or colors, n, x, y, w, h, dir h|v, gap, hex true, labels [..]. NAME.s{i}/.h{i}"""
    base = sh.get('name')
    cols = list(sh.get('colors') or sample_palette(spec.path(sh['image']), int(sh.get('n', 5)), vivid=bool(sh.get('vivid'))))
    n = len(cols)
    gap = float(sh.get('gap', 0.08))
    vertical = sh.get('dir', 'h') == 'v'
    role = _role_or(spec, 'caption', 'body')
    out = []
    for i, c in enumerate(cols):
        if vertical:
            ch = (sh['h'] - gap * (n - 1)) / n
            x, y, w, h = sh['x'], sh['y'] + i * (ch + gap), sh['w'], ch
        else:
            cw = (sh['w'] - gap * (n - 1)) / n
            x, y, w, h = sh['x'] + i * (cw + gap), sh['y'], cw, sh['h']
        out.append({'type': 'rect', 'name': _nm(base, f's{i}'), 'x': round(x, 3), 'y': round(y, 3), 'w': round(w, 3), 'h': round(h, 3),
                    'fill': c, 'line_color': sh.get('line', 'surface'), 'line_width': sh.get('line_width', 1.0),
                    **({'shadow': True} if sh.get('shadow') else {})})
        lab = (sh.get('labels') or [None] * n)[i] if i < len(sh.get('labels') or []) else None
        txt = c.upper() if sh.get('hex', True) else ''
        txt = f'{lab}\n{txt}' if lab and txt else (lab or txt)
        if txt:
            # the code sits INSIDE its own chip, in the chip's contrasting ink: never loose on a photo
            t = {'type': 'text', 'name': _nm(base, f'h{i}'), 'role': role, 'text': txt,
                 'x': round(x + 0.06, 3), 'y': round(y + 0.04, 3), 'w': round(max(0.3, w - 0.12), 3), 'h': round(max(0.2, h - 0.1), 3),
                 'color': _ink_on('#' + spec.color(c)), 'anchor': 'b'}
            if sh.get('size'):
                t['size'] = sh['size']
            out.append(t)
    return out


def _grid_pack(items, cols):
    """Row-major first fit of (span_c, span_r) items on a grid `cols` wide: [(col, row)], rows used."""
    occ, pos = set(), []
    for sc, sr in items:
        sc = min(sc, cols)
        r = 0
        while True:
            spot = next((c for c in range(cols - sc + 1)
                         if all((c + dc, r + dr) not in occ for dc in range(sc) for dr in range(sr))), None)
            if spot is not None:
                break
            r += 1
        for dc in range(sc):
            for dr in range(sr):
                occ.add((spot + dc, r + dr))
        pos.append((spot, r))
    rows = 1 + max((r for _, r in occ), default=0)
    return pos, rows


def comp_photogrid(spec, sh):
    """A dense multi-frame photo grid: bordered prints with a shadow, optional labels on colour chips.
    images [path | {path, focus, span [c, r], label}], x, y, w, h, cols, rows, gap, border, frame, tilt [deg...].
    NAME.f{i} (frame) /.p{i} (photo) /.t{i}k + .t{i} (label chip)"""
    base = sh.get('name')
    imgs = [i if isinstance(i, dict) else {'path': i} for i in sh['images']]
    cols = int(sh.get('cols', max(2, round(len(imgs) ** 0.5 + 0.49))))
    pos, rows = _grid_pack([tuple(i.get('span', [1, 1])) for i in imgs], cols)
    rows = int(sh.get('rows', rows))
    gap, bd = float(sh.get('gap', 0.12)), float(sh.get('border', 0.06))
    cw = (sh['w'] - gap * (cols - 1)) / cols
    rh = (sh['h'] - gap * (rows - 1)) / rows
    tilt = sh.get('tilt') or []
    role = _role_or(spec, sh.get('label_role', 'label'), 'label')
    out = []
    for i, (im, (c, r)) in enumerate(zip(imgs, pos)):
        sc, sr = im.get('span', [1, 1])
        x, y = sh['x'] + c * (cw + gap), sh['y'] + r * (rh + gap)
        w, h = sc * cw + (sc - 1) * gap, sr * rh + (sr - 1) * gap
        rot = tilt[i % len(tilt)] if tilt else 0
        frame = {'type': 'rect', 'name': _nm(base, f'f{i}'), 'x': round(x, 3), 'y': round(y, 3), 'w': round(w, 3), 'h': round(h, 3),
                 'fill': sh.get('frame', 'surface'), 'shadow': sh.get('shadow', True)}
        photo = {'type': 'image', 'name': _nm(base, f'p{i}'), 'path': im['path'], 'x': round(x + bd, 3), 'y': round(y + bd, 3),
                 'w': round(w - 2 * bd, 3), 'h': round(h - 2 * bd, 3), 'fit': 'cover', 'focus': im.get('focus', [0.5, 0.5])}
        for k in ('duotone', 'grayscale', 'brightness', 'contrast'):
            if k in im:
                photo[k] = im[k]
        if rot:
            frame['rot'] = photo['rot'] = rot
        out += [frame, photo]
        if im.get('label'):
            out += _chip_text(spec, base, f't{i}', im['label'], round(x + bd + 0.14, 3), round(y + h - bd - 0.36, 3), role,
                              sh.get('label_color', 'surface'), sh.get('chip', 'accent'))
    return out


def comp_nodemap(spec, sh):
    """A map of red nodes joined by lines, each label on a chip. nodes [{x, y (0..1 of the box), label, big}],
    links [[a, b], ...] (indexes or labels), x, y, w, h, image (optional map underlay), color, side r|l|t|b.
    NAME.map / .e{k} (links) / .n{i} (nodes) / .t{i}k + .t{i} (labels)"""
    base = sh.get('name')
    col = sh.get('color', 'accent')
    nodes = sh['nodes']
    idx = {n.get('id', n.get('label')): i for i, n in enumerate(nodes)}
    pt = [(sh['x'] + n['x'] * sh['w'], sh['y'] + n['y'] * sh['h']) for n in nodes]
    out = []
    if sh.get('image'):
        out.append({'type': 'image', 'name': _nm(base, 'map'), 'path': sh['image'], 'x': sh['x'], 'y': sh['y'], 'w': sh['w'], 'h': sh['h'],
                    'fit': 'cover', **{k: sh[k] for k in ('duotone', 'grayscale', 'brightness') if k in sh}})
    for k, (a, b) in enumerate(sh.get('links', [])):
        ia, ib = (idx[a] if not isinstance(a, int) else a), (idx[b] if not isinstance(b, int) else b)
        out.append({'type': 'line', 'name': _nm(base, f'e{k}'), 'x1': round(pt[ia][0], 3), 'y1': round(pt[ia][1], 3),
                    'x2': round(pt[ib][0], 3), 'y2': round(pt[ib][1], 3), 'color': col, 'width': sh.get('line_width', 1.75),
                    **({'dash': sh['dash']} if sh.get('dash') else {})})
    role = _role_or(spec, sh.get('label_role', 'label'), 'label')
    for i, (n, (px, py)) in enumerate(zip(nodes, pt)):
        r = float(n.get('r', 0.16 if n.get('big') else 0.1))
        out.append({'type': 'ellipse', 'name': _nm(base, f'n{i}'), 'x': round(px - r, 3), 'y': round(py - r, 3), 'w': round(2 * r, 3),
                    'h': round(2 * r, 3), 'fill': col, 'line_color': sh.get('node_line', 'surface'), 'line_width': 2})
        if n.get('label'):
            side = n.get('side', sh.get('side', 'r'))
            lx, ly = {'r': (px + r + 0.16, py - 0.12), 'l': (px - r - 1.6, py - 0.12),
                      't': (px - 0.5, py - r - 0.45), 'b': (px - 0.5, py + r + 0.14)}[side]
            parts = _chip_text(spec, base, f't{i}', n['label'], round(lx, 3), round(ly, 3), role,
                               sh.get('label_color', 'surface'), sh.get('chip', 'veil'))
            if n.get('desc'):
                band = parts[0]
                d = {'type': 'text', 'name': _nm(base, f'd{i}'), 'role': _role_or(spec, 'caption', 'body'), 'text': n['desc'],
                     'x': parts[1]['x'], 'y': round(band['y'] + band['h'] + 0.05, 3), 'w': float(n.get('desc_w', 2.2)), 'h': 0.4,
                     'color': sh.get('desc_color', 'ink')}
                d['h'] = round(text_height(spec, d, d['w']) + 0.04, 3)
                parts.append(d)
            if side == 'l':   # right-align the chip to the node
                dx = (px - r - 0.16) - (parts[0]['x'] + parts[0]['w'])
                for q in parts:
                    q['x'] = round(q['x'] + dx, 3)
            out += parts
    return out


EXPAND = {'heading': comp_heading, 'card': comp_card, 'stat': comp_stat, 'screen': comp_screen,
          'timeline': comp_timeline, 'bars': comp_bars, 'list': comp_list, 'quote': comp_quote,
          'callout': comp_callout, 'swatch': comp_swatch, 'photogrid': comp_photogrid, 'nodemap': comp_nodemap,
          'construct': lambda spec, sh: __import__('construct').comp_construct(spec, sh)}   # v5.5: construction motion


def expand(spec, sh):
    """Primitive shapes for one spec shape (components expand; groups recurse)."""
    if sh.get('hidden'):
        return []
    if sh.get('anim') and sh.get('type') in EXPAND:
        sh = dict(sh)
        anim = sh.pop('anim')
        parts = expand(spec, sh)
        step = int(sh.get('stagger', 120))
        anims = anim if isinstance(anim, list) else [anim]
        for i, p in enumerate(parts):
            p['anim'] = [{**a, 'delay': int(a.get('delay', 0)) + i * step} for a in anims]
        return parts
    sh = dict(sh)
    if 'col' in sh:
        sh['x'], sh['w'] = spec.col(*sh['col'])
    t = sh.get('type')
    if t == 'poly' and sh.get('points'):
        xs = [p[0] for p in sh['points']]
        ys = [p[1] for p in sh['points']]
        sh.update({'x': min(xs), 'y': min(ys), 'w': max(xs) - min(xs), 'h': max(ys) - min(ys)})
    if t in EXPAND:
        for k in ('x', 'y') if t == 'callout' else ('x', 'y', 'w'):   # a callout chip sizes itself to its label
            if k not in sh:
                raise SystemExit(f'{t} {sh.get("name")!r}: needs {k} (or "col")')
        prims = []
        for p in EXPAND[t](spec, sh):
            prims.extend(expand(spec, p))
        return prims
    if t == 'group':
        kids = []
        for k in sh['shapes']:
            kids.extend(expand(spec, k))
        sh['shapes'] = kids
        if 'x' not in sh and kids:  # bounding box for the layout checks
            boxes = []
            for k in kids:
                if 'x1' in k:
                    boxes.append((min(k['x1'], k['x2']), min(k['y1'], k['y2']), max(k['x1'], k['x2']), max(k['y1'], k['y2'])))
                elif 'x' in k and 'w' in k:
                    boxes.append((k['x'], k['y'], k['x'] + k['w'], k['y'] + k.get('h', 0)))
            if boxes:
                x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
                sh.update({'x': x0, 'y': y0, 'w': max(b[2] for b in boxes) - x0, 'h': max(b[3] for b in boxes) - y0})
    return [sh]


# ---------------------------------------------------------------- XML helpers

def _spPr_insert(spPr, el):
    """Insert an effect element in schema order (after ln, else fill, else geometry, else xfrm)."""
    for tag in ('a:ln',):
        found = spPr.find(qn(tag))
        if found is not None:
            found.addnext(el)
            return
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        found = spPr.find(qn(tag))
        if found is not None:
            found.addnext(el)
            return
    for tag in ('a:prstGeom', 'a:custGeom', 'a:xfrm'):
        found = spPr.find(qn(tag))
        if found is not None:
            found.addnext(el)
            return
    spPr.append(el)


def apply_shadow(spPr, sd, spec):
    if not sd:
        return
    sd = {**spec.shadow, **(sd if isinstance(sd, dict) else {})}
    for old in spPr.findall(qn('a:effectLst')):
        spPr.remove(old)
    eff = etree.Element(qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad=str(int(sd['blur'] * 12700)), dist=str(int(sd['dist'] * 12700)),
                          dir=str(int(sd['dir'] * 60000)), algn='ctr', rotWithShape='0')
    c = etree.SubElement(sh, qn('a:srgbClr'), val=spec.color(sd['color']))
    etree.SubElement(c, qn('a:alpha'), val=str(int(sd['alpha'] * 100000)))
    _spPr_insert(spPr, eff)


def set_alpha_fill(shape, hexcol, alpha):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexcol)
    if alpha is not None and alpha < 1:
        clr = shape.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(clr, qn('a:alpha'), val=str(int(alpha * 100000)))


def set_gradient(shape, stops, angle, spec):
    spPr = shape.fill._xPr
    for tag in ('a:solidFill', 'a:noFill', 'a:gradFill', 'a:blipFill', 'a:pattFill'):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    grad = etree.Element(qn('a:gradFill'), rotWithShape='1')
    lst = etree.SubElement(grad, qn('a:gsLst'))
    for st in stops:  # [colour, alpha 0..1, position 0..100]
        col, alpha, pos = st[0], st[1], st[2]
        gs = etree.SubElement(lst, qn('a:gs'), pos=str(int(pos * 1000)))
        c = etree.SubElement(gs, qn('a:srgbClr'), val=spec.color(col))
        etree.SubElement(c, qn('a:alpha'), val=str(int(alpha * 100000)))
    etree.SubElement(grad, qn('a:lin'), ang=str(int(angle * 60000)), scaled='0')
    spPr.find(qn('a:prstGeom')).addnext(grad)


def line_alpha(line_el, alpha):
    if alpha is None or alpha >= 1:
        return
    clr = line_el.find(qn('a:solidFill'))
    if clr is not None and clr.find(qn('a:srgbClr')) is not None:
        etree.SubElement(clr.find(qn('a:srgbClr')), qn('a:alpha'), val=str(int(alpha * 100000)))


# ---------------------------------------------------------------- primitives

def add_box(shapes, sh, spec, kind):
    if kind == 'ellipse':
        geo = MSO_SHAPE.OVAL
    elif sh.get('geom'):
        try:
            geo = getattr(MSO_SHAPE, sh['geom'].upper())
        except AttributeError:
            raise SystemExit(f'unknown geom {sh["geom"]!r} (use an MSO_SHAPE name, e.g. CHEVRON, TRIANGLE, ROUNDED_RECTANGLE)')
    elif sh.get('radius'):
        geo = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        geo = MSO_SHAPE.RECTANGLE
    shp = shapes.add_shape(geo, Inches(sh['x']), Inches(sh['y']), Inches(max(0.001, sh['w'])), Inches(max(0.001, sh['h'])))
    if geo == MSO_SHAPE.ROUNDED_RECTANGLE and sh.get('radius'):
        shp.adjustments[0] = min(0.5, sh['radius'])
    if sh.get('gradient'):
        set_gradient(shp, sh['gradient'], sh.get('angle', 90), spec)
    elif sh.get('fill'):
        set_alpha_fill(shp, spec.color(sh['fill']), sh.get('alpha'))
    else:
        shp.fill.background()
    if sh.get('line_color'):
        shp.line.color.rgb = RGBColor.from_string(spec.color(sh['line_color']))
        shp.line.width = Pt(sh.get('line_width', 1))
        line_alpha(shp.line._get_or_add_ln(), sh.get('line_alpha'))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if sh.get('shadow'):
        apply_shadow(shp._element.spPr, sh['shadow'], spec)
    if shp.has_text_frame:
        shp.text_frame.text = ''
    return shp


def add_line(shapes, sh, spec):
    ln = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(sh['x1']), Inches(sh['y1']), Inches(sh['x2']), Inches(sh['y2']))
    ln.line.color.rgb = RGBColor.from_string(spec.color(sh.get('color', '#FFFFFF')))
    ln.line.width = Pt(sh.get('width', 1))
    lnel = ln.line._get_or_add_ln()
    line_alpha(lnel, sh.get('alpha'))
    if sh.get('cap'):
        lnel.set('cap', {'round': 'rnd', 'flat': 'flat', 'square': 'sq'}[sh['cap']])
    if sh.get('dash'):
        etree.SubElement(lnel, qn('a:prstDash'), val={'dash': 'dash', 'dot': 'sysDot', 'longdash': 'lgDash'}[sh['dash']])
    for end, key in (('a:headEnd', 'arrow_start'), ('a:tailEnd', 'arrow')):
        if sh.get(key):  # "arrow": true | "triangle" | "stealth" | "oval" (head at x2,y2)
            kind = sh[key] if isinstance(sh[key], str) else 'triangle'
            etree.SubElement(lnel, qn(end), type=kind, w='med', len='med')
    if sh.get('shadow'):
        apply_shadow(ln._element.spPr, sh['shadow'], spec)
    return ln


def add_poly(shapes, sh, spec):
    """A freeform polygon from inch points (isometric faces, paper-cut layers, masks, arrows)."""
    pts = sh['points']
    if len(pts) < 2:
        raise SystemExit(f'poly {sh.get("name")!r}: needs at least 2 points')
    fb = shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=poly_closed(sh))
    shp = fb.convert_to_shape()
    if sh.get('gradient'):
        set_gradient_any(shp, sh['gradient'], sh.get('angle', 90), spec)
    elif sh.get('fill'):
        set_alpha_fill(shp, spec.color(sh['fill']), sh.get('alpha'))
    else:
        shp.fill.background()
    if sh.get('line_color'):
        shp.line.color.rgb = RGBColor.from_string(spec.color(sh['line_color']))
        shp.line.width = Pt(sh.get('line_width', 1))
        line_alpha(shp.line._get_or_add_ln(), sh.get('line_alpha'))
        if sh.get('cap'):
            shp.line._get_or_add_ln().set('cap', {'round': 'rnd', 'flat': 'flat', 'square': 'sq'}[sh['cap']])
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if sh.get('shadow'):
        apply_shadow(shp._element.spPr, sh['shadow'], spec)
    return shp


def set_gradient_any(shape, stops, angle, spec):
    """set_gradient for shapes whose geometry is custom (freeform)."""
    spPr = shape._element.spPr
    for tag in ('a:solidFill', 'a:noFill', 'a:gradFill', 'a:blipFill', 'a:pattFill'):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    grad = etree.Element(qn('a:gradFill'), rotWithShape='1')
    lst = etree.SubElement(grad, qn('a:gsLst'))
    for st in stops:
        gs = etree.SubElement(lst, qn('a:gs'), pos=str(int(st[2] * 1000)))
        c = etree.SubElement(gs, qn('a:srgbClr'), val=spec.color(st[0]))
        etree.SubElement(c, qn('a:alpha'), val=str(int(st[1] * 100000)))
    etree.SubElement(grad, qn('a:lin'), ang=str(int(angle * 60000)), scaled='0')
    geo = spPr.find(qn('a:custGeom'))
    if geo is None:
        geo = spPr.find(qn('a:prstGeom'))
    geo.addnext(grad)


def add_video(shapes, sh, spec):
    """An embedded MP4 (looping background, a rendered Manim/diagram scene). The poster frame is what
    LibreOffice, Google Slides import and printouts show; PowerPoint plays the video."""
    p = spec.path(sh['path'])
    if not p.exists():
        raise SystemExit(f'video not found: {p}')
    poster = spec.path(sh['poster']) if sh.get('poster') else p.with_suffix('.poster.png')
    if not poster.exists():
        # v5.7.1: a missing poster (manim films have no poster step) is grabbed from the film instead of a
        # python-pptx FileNotFoundError: the LAST frame by default (a film's end state is its editable picture),
        # "poster_at": seconds for another one (film_check.grab_frame copes with manimgx's variable frame rate).
        from film_check import grab_frame
        at = sh.get('poster_at')
        if not grab_frame(p, poster, at):
            raise SystemExit(f'video {p.name}: no poster at {poster} and ffmpeg could not grab one')
        spec.warnings.append(f'video {p.name}: poster {poster.name} was missing; grabbed the '
                             f'{"frame at " + str(at) + " s" if at is not None else "last frame"} (look at it)')
    mov = shapes.add_movie(str(p), Inches(sh['x']), Inches(sh['y']), Inches(sh['w']), Inches(sh['h']),
                           poster_frame_image=str(poster), mime_type='video/mp4')
    return mov


_SAFE_CACHE = {}


def ppt_safe(p):
    """v5.3: MPO (phone/3D cameras), EXIF-rotated, CMYK, TIFF or WebP files become a cached RGB JPEG/PNG
    so python-pptx accepts them and PowerPoint shows them the right way up."""
    p = Path(p)
    if p in _SAFE_CACHE:
        return _SAFE_CACHE[p]
    out = p
    try:
        with Image.open(p) as im:
            fmt = (im.format or '').upper()
            orient = im.getexif().get(0x0112, 1)
            if fmt in ('MPO', 'TIFF', 'WEBP', 'BMP') or orient not in (None, 1) or im.mode == 'CMYK':
                from PIL import ImageOps
                import hashlib
                import tempfile
                im.seek(0)
                fixed = ImageOps.exif_transpose(im)
                alpha = fixed.mode in ('RGBA', 'LA')
                key = hashlib.sha1(f'{p.resolve()}:{p.stat().st_mtime_ns}'.encode()).hexdigest()[:12]
                out = Path(tempfile.gettempdir()) / f'pptsafe-{key}{".png" if alpha else ".jpg"}'
                if not out.exists():
                    (fixed if alpha else fixed.convert('RGB')).save(out, quality=92)
    except OSError:
        out = p
    _SAFE_CACHE[p] = out
    return out


def add_image(shapes, sh, spec):
    p = spec.path(sh['path'])
    if not p.exists():
        raise SystemExit(f'image not found: {p}')
    p = ppt_safe(p)
    pic = shapes.add_picture(str(p), Inches(sh['x']), Inches(sh['y']), Inches(sh['w']), Inches(sh['h']))
    fit = sh.get('fit', 'cover')
    iw, ih = Image.open(p).size
    box = sh['w'] / sh['h']
    img = iw / ih
    fx, fy = sh.get('focus', [0.5, 0.5])
    if fit == 'cover' and abs(img - box) > 0.005:
        if img > box:
            keep = box / img
            left = (1 - keep) * fx
            pic.crop_left, pic.crop_right = left, 1 - keep - left
        else:
            keep = img / box
            top = (1 - keep) * fy
            pic.crop_top, pic.crop_bottom = top, 1 - keep - top
    elif fit == 'contain' and abs(img - box) > 0.005:
        if img > box:
            nh = sh['w'] / img
            pic.top = Inches(sh['y'] + (sh['h'] - nh) / 2)
            pic.height = Inches(nh)
        else:
            nw = sh['h'] * img
            pic.left = Inches(sh['x'] + (sh['w'] - nw) / 2)
            pic.width = Inches(nw)
    blip = pic._element.find('.//' + qn('a:blip'))
    if sh.get('alpha') is not None and sh['alpha'] < 1:
        etree.SubElement(blip, qn('a:alphaModFix'), amt=str(int(sh['alpha'] * 100000)))
    if sh.get('grayscale'):
        etree.SubElement(blip, qn('a:grayscl'))
    if sh.get('duotone'):
        duo = etree.SubElement(blip, qn('a:duotone'))
        for c in sh['duotone'][:2]:
            etree.SubElement(duo, qn('a:srgbClr'), val=spec.color(c))
    if sh.get('brightness') or sh.get('contrast'):
        etree.SubElement(blip, qn('a:lum'), bright=str(int(sh.get('brightness', 0) * 100000)),
                         contrast=str(int(sh.get('contrast', 0) * 100000)))
    if sh.get('radius'):
        geom = pic._element.spPr.find(qn('a:prstGeom'))
        geom.set('prst', 'roundRect')
        av = geom.find(qn('a:avLst'))
        if av is None:
            av = etree.SubElement(geom, qn('a:avLst'))
        for g in list(av):
            av.remove(g)
        etree.SubElement(av, qn('a:gd'), name='adj', fmla=f'val {int(min(0.5, sh["radius"]) * 100000)}')
    if sh.get('line_color'):
        pic.line.color.rgb = RGBColor.from_string(spec.color(sh['line_color']))
        pic.line.width = Pt(sh.get('line_width', 1))
    if sh.get('shadow'):
        apply_shadow(pic._element.spPr, sh['shadow'], spec)
    return pic


def style_run(run, d, spec):
    f = run.font
    f.name = spec.font(d['font'])
    f.size = Pt(d['size'])
    f.bold = d['bold']
    f.italic = d['italic']
    col = d['color']
    if d.get('fade') is not None and d['fade'] < 1:
        # A quiet text is a solid colour mixed toward its background, not a transparent run:
        # LibreOffice misplaces Vietnamese marks on transparent text, PowerPoint does not care.
        col = mix(spec, col, d.get('fade_bg') or 'ground', d['fade'])
    f.color.rgb = RGBColor.from_string(spec.color(col))
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):  # same family for every script slot (Vietnamese renders via latin)
        e = rPr.find(qn(tag))
        if e is None:
            e = etree.SubElement(rPr, qn(tag))
        e.set('typeface', f.name)
    if d['spacing']:
        rPr.set('spc', str(int(d['spacing'] * 100)))
    if d['caps']:
        rPr.set('cap', 'all')
    a = d.get('alpha')
    if a is not None and a < 1:
        clr = rPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(clr, qn('a:alpha'), val=str(int(a * 100000)))


def add_text(shapes, sh, spec):
    tb = shapes.add_textbox(Inches(sh['x']), Inches(sh['y']), Inches(sh['w']), Inches(sh['h']))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, 0)
    base = text_props(spec, sh)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[base['anchor']]
    paras = sh['text'] if isinstance(sh['text'], list) else [sh['text']]
    for i, para in enumerate(paras):
        d = para if isinstance(para, dict) else {'text': para}
        pd = text_props(spec, {**sh, **{k: v for k, v in d.items() if k not in ('text', 'runs')}})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[pd['align']]
        if pd['line'] and abs(pd['line'] - 1.0) > 1e-6:
            p.line_spacing = pd['line']
        if d.get('space_before'):
            p.space_before = Pt(d['space_before'])
        for rd in d.get('runs') or [d]:
            r = p.add_run()
            r.text = str(rd.get('text', ''))
            rp = text_props(spec, {**sh, **{k: v for k, v in d.items() if k not in ('text', 'runs')},
                                   **{k: v for k, v in rd.items() if k != 'text'}})
            style_run(r, rp, spec)
    return tb


def add_shape(shapes, sh, spec):
    t = sh.get('type')
    if t == 'image':
        obj = add_image(shapes, sh, spec)
    elif t == 'text':
        obj = add_text(shapes, sh, spec)
    elif t in ('rect', 'ellipse'):
        obj = add_box(shapes, sh, spec, t)
    elif t == 'line':
        obj = add_line(shapes, sh, spec)
    elif t == 'poly':
        obj = add_poly(shapes, sh, spec)
    elif t == 'video':
        obj = add_video(shapes, sh, spec)
    elif t == 'group':
        obj = shapes.add_group_shape()
        obj.shapes._recalculate_extents = lambda: None   # once at the end, not after every child
        for k in sh['shapes']:
            add_shape(obj.shapes, k, spec)
        obj._element.recalculate_extents()
        if sh.get('shadow'):
            apply_shadow(obj._element.grpSpPr, sh['shadow'], spec)
    else:
        raise SystemExit(f'unknown shape type {t!r} (primitives: text image rect ellipse line poly video group; components: {", ".join(COMPONENTS)})')
    if sh.get('name'):
        obj.name = sh['name']
    rot = sh.get('rot', spec.role(sh.get('role')).get('rot') if t == 'text' and sh.get('role') else None)
    if rot:
        obj.rotation = rot
    return obj


# ---------------------------------------------------------------- slide assembly

def set_background(slide, spec, bg):
    if bg is None:
        return
    img = None
    if isinstance(bg, dict):
        img = spec.path(bg['image'])
    elif isinstance(bg, str):
        img = spec.ground(bg)
    if img is None:
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = RGBColor.from_string(spec.color(bg))
        return
    img = ppt_safe(img)
    _, rId = slide.part.get_or_add_image_part(str(img))
    iw, ih = Image.open(img).size
    want = W / H
    l = r = t = b = 0
    if iw / ih > want + 0.005:
        keep = want / (iw / ih)
        l = r = int((1 - keep) / 2 * 100000)
    elif iw / ih < want - 0.005:
        keep = (iw / ih) / want
        t = b = int((1 - keep) / 2 * 100000)
    cSld = slide._element.cSld
    old = cSld.find(qn('p:bg'))
    if old is not None:
        cSld.remove(old)
    xml = (f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
           f'<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rId}"/>'
           f'<a:srcRect l="{l}" t="{t}" r="{r}" b="{b}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
           f'<a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, etree.fromstring(xml))


def chrome_for(spec, s, n, total):
    keep = s.get('chrome', True)
    over = s.get('chrome_set', {})
    # v5.7.1: {scene} is the audience-facing chapter name ("scene_title" on the slide, or deck "scene_titles":
    # {id: name}); a bare internal id such as "S4" is never printed on a slide
    sid = str(s.get('scene', '') or '')
    scene_txt = s.get('scene_title') or (getattr(spec, 'raw', {}) or {}).get('scene_titles', {}).get(sid) or \
        ('' if re.fullmatch(r'[A-Za-z]{0,3}\d+[a-z]?(?:[._-]\w+)?', sid) else sid)
    tokens = {'n': f'{n:02d}', 'total': f'{total:02d}', 'title': spec.title, 'source': s.get('source', ''),
              'scene': scene_txt}
    layers = {}
    for layer in ('under', 'over'):
        out = []
        for sh in spec.chrome.get(layer, []):
            name = sh.get('name')
            if keep is False or (isinstance(keep, list) and name not in keep):
                continue
            sh = merge(sh, over.get(name, {}))
            if sh.get('type') == 'text':
                txt = sh['text'] if isinstance(sh['text'], str) else sh['text']
                if isinstance(txt, str):
                    txt = re.sub(r'\{(\w+)\}', lambda m: str(tokens.get(m.group(1), m.group(0))), txt)
                    txt = re.sub(r'^(?:\s*[·|/–—-]\s*)+|(?:\s*[·|/–—-]\s*)+$', '', txt)   # a separator left by an empty token
                    if not txt.strip():
                        continue
                    sh['text'] = txt
            out.append(sh)
        layers[layer] = out
    return layers


def check_text(spec, sh, where, chrome_names):
    """Warnings for text that overflows its box or leaves the safe area."""
    if sh.get('type') == 'group':
        for k in sh['shapes']:
            check_text(spec, k, where, chrome_names)
        return
    if sh.get('type') != 'text':
        return
    name = sh.get('name') or '(unnamed text)'
    if name in chrome_names or sh.get('bleed'):
        return
    g = spec.grid
    m = g['safe']
    bx, by, bw, bh = sh['x'], sh['y'], sh['w'], sh['h']
    if sh.get('rot'):   # v5.2: a rotated text box is judged by the box it really covers on screen
        import math as _m
        r = _m.radians(sh['rot'])
        rw, rh = abs(bw * _m.cos(r)) + abs(bh * _m.sin(r)), abs(bw * _m.sin(r)) + abs(bh * _m.cos(r))
        bx, by, bw, bh = bx + bw / 2 - rw / 2, by + bh / 2 - rh / 2, rw, rh
    if bx < m - 1e-3 or by < m - 1e-3 or bx + bw > W - m + 1e-3 or by + bh > H - m + 1e-3:
        if not (bx > W or by > H or bx + bw < 0 or by + bh < 0):  # parked off-slide is fine
            spec.warnings.append(f'{where} {name}: outside the safe area ({m:.2f} in from every edge)')
    import typeset
    paras = sh['text'] if isinstance(sh['text'], list) else [sh['text']]
    need, problems, missing = 0.0, [], None
    for para in paras:
        d = para if isinstance(para, dict) else {'text': para}
        p = text_props(spec, {**sh, **{k: v for k, v in d.items() if k not in ('text', 'runs')}})
        txt = d.get('text') if 'runs' not in d else ''.join(r.get('text', '') for r in d['runs'])
        txt = str(txt or ' ')
        mm = typeset.measure(txt, p['font'], p['size'], sh['w'], p['bold'], p['italic'], p['line'], p['spacing'], p['caps'])
        if mm['font_missing']:
            missing = p['font']
            break
        need += mm['height_in'] + d.get('space_before', 0) / 72
        if mm['longest_word_in'] > sh['w'] * 1.03:
            problems.append(f'a word is wider than the box ({mm["longest_word_in"]:.2f} in > {sh["w"]:.2f} in)')
        if len(mm['lines']) >= 2 and len(mm['lines'][-1].split()) == 1 and len(txt.split()) >= 4 and p['size'] >= 28:
            problems.append(f'a single word on the last line ("{mm["lines"][-1]}")')
    if missing:
        key = f'font not installed: {missing}'
        if key not in spec.warnings:
            spec.warnings.append(key)
        return
    if need > sh['h'] * 1.06 + 0.03:
        problems.append(f'{need:.2f} in of text in a {sh["h"]:.2f} in box: the lines need more height')
    for pr in problems:
        spec.warnings.append(f'{where} {name}: {pr}')


def text_extent(spec, sh):
    """(widest line, total height) in inches of a text shape laid out in its box."""
    import typeset
    paras = sh['text'] if isinstance(sh['text'], list) else [sh['text']]
    wide, total = 0.0, 0.0
    for para in paras:
        d = para if isinstance(para, dict) else {'text': para}
        p = text_props(spec, {**sh, **{k: v for k, v in d.items() if k not in ('text', 'runs')}})
        txt = d.get('text') if 'runs' not in d else ''.join(r.get('text', '') for r in d['runs'])
        m = typeset.measure(str(txt or ' '), p['font'], p['size'], sh['w'], p['bold'], p['italic'], p['line'], p['spacing'], p['caps'])
        if m['font_missing']:
            return sh['w'], text_height(spec, sh, sh['w'])
        wide = max(wide, min(sh['w'], m['widest_in']))
        total += m['height_in'] + d.get('space_before', 0) / 72
    return wide, total


def _boxes(spec, shapes, texts, groups, owner=None):
    for sh in shapes:
        if sh.get('type') == 'group':
            kids = sh['shapes']
            xs = [k.get('x', min(k.get('x1', 0), k.get('x2', 0))) for k in kids]
            ys = [k.get('y', min(k.get('y1', 0), k.get('y2', 0))) for k in kids]
            x2 = [k.get('x', 0) + k.get('w', 0) if 'w' in k else max(k.get('x1', 0), k.get('x2', 0)) for k in kids]
            y2 = [k.get('y', 0) + k.get('h', 0) if 'h' in k else max(k.get('y1', 0), k.get('y2', 0)) for k in kids]
            if kids:
                groups.append((sh.get('name') or '(group)', min(xs), min(ys), max(x2) - min(xs), max(y2) - min(ys)))
            _boxes(spec, kids, texts, groups, owner=sh.get('name') or '(group)')
        elif sh.get('type') == 'text' and str(sh.get('text', '')).strip():
            wide, need = text_extent(spec, sh)
            tp = text_props(spec, sh)
            anchor, align = tp['anchor'], tp['align']
            top = sh['y'] if anchor == 't' else (sh['y'] + sh['h'] - need if anchor == 'b' else sh['y'] + (sh['h'] - need) / 2)
            left = sh['x'] if align == 'l' else (sh['x'] + sh['w'] - wide if align == 'r' else sh['x'] + (sh['w'] - wide) / 2)
            rot = sh.get('rot', spec.role(sh.get('role')).get('rot'))
            if rot:   # v5.2: rotate the measured ink box about the shape centre
                import math as _m
                r = _m.radians(rot)
                cx0, cy0 = sh['x'] + sh['w'] / 2, sh['y'] + sh['h'] / 2
                pts = [(left + dx - cx0, top + dy - cy0) for dx in (0, wide) for dy in (0, need)]
                pts = [(cx0 + px * _m.cos(r) - py * _m.sin(r), cy0 + px * _m.sin(r) + py * _m.cos(r)) for px, py in pts]
                left, top = min(p[0] for p in pts), min(p[1] for p in pts)
                wide, need = max(p[0] for p in pts) - left, max(p[1] for p in pts) - top
            overlay = bool(sh.get('overlay', spec.role(sh.get('role')).get('overlay', False)))   # v5.4: a script accent may cross a title
            texts.append((sh.get('name') or '(text)', left, top, wide, need, owner, overlay))


def _overlap(a, b):
    ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
    return ox, oy


def check_collisions(spec, prims, where, chrome_names):
    """Texts that overlap each other, or a text that runs over a card or screen it does not
    belong to, are almost always layout defects (a title that grew into the cards below it).
    Measured text extents are compared, not box sizes."""
    texts, groups = [], []
    _boxes(spec, prims, texts, groups)
    texts = [t for t in texts if t[0] not in chrome_names]
    seen = set()
    for i, a in enumerate(texts):
        for b in texts[i + 1:]:
            if a[6] or b[6]:
                continue
            ox, oy = _overlap(a, b)
            if ox > 0.05 and oy > 0.04 and (a[0], b[0]) not in seen:
                seen.add((a[0], b[0]))
                spec.warnings.append(f'{where}: {a[0]} and {b[0]} overlap on screen ({oy:.2f} in): '
                                     'move one, shorten the text or widen the column')
        for g in groups:
            if a[5] == g[0]:
                continue
            ox, oy = _overlap(a, g)
            inside = a[1] >= g[1] - 0.01 and a[2] >= g[2] - 0.01 and a[1] + a[3] <= g[1] + g[3] + 0.01 and a[2] + a[4] <= g[2] + g[4] + 0.01
            if ox > 0.05 and oy > 0.04 and not inside:
                spec.warnings.append(f'{where}: {a[0]} runs over {g[0]} ({oy:.2f} in): give the text its own space')


def build(data, root, base=None):
    spec = Spec(data, root)
    prs = Presentation(base) if base else Presentation()
    if not base:
        prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    slides = [s for s in resolve(data['slides']) if not s.get('hidden')]   # v6.22: slide-level "hidden": true = CUT (kept in deck.json)
    total = len(slides)
    made, links = {}, []
    for n, s in enumerate(slides, 1):
        sl = prs.slides.add_slide(layout)
        if s.get('id'):
            made[s['id']] = sl
        deck_spec, spec = spec, spec.for_style(s.get('style'))
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        set_background(sl, spec, s.get('bg', spec.default_bg))
        bgp = isinstance(s.get('bg', spec.default_bg), dict) and bool((s.get('bg') or {}).get('image'))
        if s.get('scene') or s.get('bookend') or s.get('still') or bgp:
            # v5.4: "bookend": true = a declared circular close (image reuse allowed);
            # "still": true = a deliberately still page (excluded from the moving-depth share)
            # v5.6 (I21): "[bg-photo]" = the slide background is a picture (not a ground texture), for variety_audit
            sl._element.cSld.set('name', (str(s.get('scene') or '') + (' [bookend]' if s.get('bookend') else '')
                                          + (' [still]' if s.get('still') else '') + (' [bg-photo]' if bgp else '')).strip())
        layers = chrome_for(spec, s, n, total)
        chrome_names = {c.get('name') for c in layers['under'] + layers['over']}
        prims = []
        for sh in layers['under'] + s['shapes'] + layers['over']:
            prims.extend(expand(spec, sh))
        names = [p.get('name') for p in prims if (p.get('name') or '').startswith('!!')]
        dup = sorted({x for x in names if names.count(x) > 1})
        if dup:
            raise SystemExit(f'slide {n}: duplicate !! names after expanding components/chrome: {dup}')
        s['_anims'], s['_media'] = [], {}
        for k, sh in enumerate(prims):
            if sh.get('anim'):
                if not sh.get('name'):
                    sh['name'] = f'fx{n}-{k}'
                for a in (sh['anim'] if isinstance(sh['anim'], list) else [sh['anim']]):
                    s['_anims'].append({'shape': sh['name'], **a})
            if sh.get('type') == 'group':
                bad = [kid.get('name') for kid in sh['shapes'] if kid.get('anim')]
                if bad:
                    spec.warnings.append(f'slide {n}: "anim" on shapes inside group {sh.get("name")!r} is ignored '
                                         f'(PowerPoint animates only top-level shapes): {bad}')
            if sh.get('type') == 'video':
                s['_media'][sh.get('name') or f'video{n}-{k}'] = {kk: sh[kk] for kk in ('loop', 'mute', 'autoplay', 'delay') if kk in sh}
                if not sh.get('name'):
                    sh['name'] = f'video{n}-{k}'
                try:
                    import subprocess
                    r = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0',
                                        str(spec.path(sh['path']))], capture_output=True, text=True)
                    s['_media'][sh['name']]['ms'] = int(float(r.stdout.strip()) * 1000)
                except (ValueError, OSError):
                    pass
            check_text(spec, sh, f'slide {n}', chrome_names)
            obj = add_shape(sl.shapes, sh, spec)
            if sh.get('link'):
                links.append((n, obj, sh['link']))
        check_collisions(spec, prims, f'slide {n}', chrome_names)
        s['_prims'] = [p.get('name') for p in prims]
        if s.get('notes'):
            sl.notes_slide.notes_text_frame.text = s['notes']
        else:
            sl.notes_slide  # an empty notes page so notes-set can fill it later
        spec = deck_spec
    for n, obj, target in links:
        if target not in made:
            raise SystemExit(f'slide {n}: link {target!r}: no slide with that id')
        obj.click_action.target_slide = made[target]
    return prs, slides, spec


def motion_config(data, slides, spec, out):
    """motion.json next to the output when the deck declares motion; WORK/motion.json values sit underneath."""
    declares = 'motion' in data or data.get('compat', 'safe') != 'morph' or any(k in s for s in slides for k in ('ms', 'morph', 'keep_timing', 'transition')) \
        or any(s.get('_anims') or s.get('_media') for s in slides)
    if not declares:
        return None
    cfg = {'default_ms': spec.motion['default_ms'], 'stagger_ms': spec.motion['stagger_ms'], 'dur': {}, 'option': {}, 'keep_timing': [],
           'transition': {}, 'effects': {}, 'media': {}}
    import compat
    try:
        cfg['compat'] = compat.check_mode(data.get('compat'))
    except ValueError as e:
        raise SystemExit(str(e))
    if data.get('safe_transition'):
        cfg['safe_transition'] = data['safe_transition']
    work = Path(out).resolve().parent.parent / 'motion.json'
    if work.is_file():
        try:
            old = json.loads(work.read_text(encoding='utf-8'))
            cfg['dur'].update({str(k): v for k, v in old.get('dur', {}).items()})
            cfg['option'].update({str(k): v for k, v in old.get('option', {}).items()})
            cfg['keep_timing'] = sorted(set(cfg['keep_timing']) | {int(x) for x in old.get('keep_timing', [])})
            for key in ('transition', 'effects', 'media'):
                cfg[key].update({str(k): v for k, v in old.get(key, {}).items()})
        except (ValueError, OSError):
            pass
    for n, s in enumerate(slides, 1):
        if s.get('ms'):
            cfg['dur'][str(n)] = int(s['ms'])
        if s.get('morph'):
            if s['morph'] not in ('byObject', 'byWord', 'byChar'):
                raise SystemExit(f'slide {n}: morph must be byObject, byWord or byChar')
            cfg['option'][str(n)] = s['morph']
        if s.get('keep_timing'):
            cfg['keep_timing'] = sorted(set(cfg['keep_timing']) | {n})
        if s.get('transition'):
            import motionlib
            t = s['transition'] if isinstance(s['transition'], dict) else {'type': s['transition']}
            if t.get('type', 'morph') not in motionlib.TRANSITIONS:
                raise SystemExit(f'slide {n}: unknown transition {t.get("type")!r}; known: {", ".join(motionlib.transition_names())}')
            cfg['transition'][str(n)] = t
        if s.get('_anims'):
            import motionlib
            for a in s['_anims']:
                if a.get('fx', 'fade') not in motionlib.PRESETS:
                    raise SystemExit(f'slide {n}: {a["shape"]}: unknown effect {a.get("fx")!r}; known: {", ".join(sorted(motionlib.PRESETS))}')
            cfg['effects'][str(n)] = s['_anims']
        if s.get('_media'):
            cfg['media'][str(n)] = s['_media']
    p = Path(out).parent / 'motion.json'
    p.write_text(json.dumps(cfg, indent=1), encoding='utf-8')
    return p


def strip_theme_style(prs):
    """v5.7: drop the theme style reference (<p:style>) from drawn shapes that carry no text.
    python-pptx autoshapes and freeforms get effectRef idx=2 from it; LibreOffice (and some PowerPoint themes)
    then paint the theme's soft shadow under every hairline, which turns a map graticule or a dense poly
    drawing into grey bands. Every drawn shape here sets fill, line and effects explicitly, so the reference
    is never wanted. Shapes with text keep it (their fontRef supplies the default text colour)."""
    n = 0
    ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    for slide in prs.slides:
        for st in list(slide.shapes._spTree.iter(ns + 'style')):
            sp = st.getparent()
            if ''.join(t for t in sp.itertext() if t.strip()):
                continue
            sp.remove(st)
            n += 1
    return n


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('spec')
    ap.add_argument('--out', required=True)
    ap.add_argument('--base', help='optional .pptx whose masters/theme to reuse (usually an empty one)')
    ap.add_argument('--quiet', action='store_true', help='only print the summary line and warnings count')
    a = ap.parse_args(argv)
    sys.path.insert(0, str(HERE))
    sp = Path(a.spec)
    data = json.loads(sp.read_text(encoding='utf-8'))
    prs, slides, spec = build(data, sp.parent.resolve(), a.base)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    nstyle = strip_theme_style(prs)
    prs.save(a.out)
    mj = motion_config(data, slides, spec, a.out)
    print(f'built {len(slides)} slides -> {a.out}' + (f' (+ {mj.name})' if mj else ''))
    seen = []
    for w in spec.warnings:
        if w not in seen:
            seen.append(w)
    for w in (seen if not a.quiet else seen[:12]):
        print('WARN', w)
    if seen:
        print(f'{len(seen)} layout warning(s): fix them now, they show up on screen')
    print('NEXT: build.py runs the rest (ooxml unpack -> motion.py -> pack -> render.sh -> film_check), then LOOK')
    return 0


if __name__ == '__main__':
    sys.exit(main())
