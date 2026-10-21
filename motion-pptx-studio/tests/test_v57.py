"""v5.7: the "Khí áp" flagship lessons: theme style refs stripped from drawn shapes (LibreOffice grey bands),
anatomy anchors labelled, the creative-authority section, the manimgx quirks reference and the flagship example."""
import json
import zipfile
from pathlib import Path

from conftest import SKILL

import build_deck


def _deck(tmp_path):
    spec = {'title': 't', 'slides': [{'id': 'a', 'bg': '#EEEBE3', 'shapes': [
        {'type': 'poly', 'name': 'grat', 'points': [[1, 1], [5, 1.2], [9, 1.1]], 'open': True, 'line_color': '#C9C3B6', 'line_width': 0.5},
        {'type': 'rect', 'name': 'box', 'x': 1, 'y': 2, 'w': 2, 'h': 1, 'fill': '#1F4F8F'},
        {'type': 'ellipse', 'name': 'dot', 'x': 4, 'y': 2, 'w': 0.3, 'h': 0.3, 'fill': '#C0391F'},
        {'type': 'text', 'name': 'lab', 'text': 'Khí áp', 'x': 1, 'y': 4, 'w': 3, 'h': 0.6, 'size': 23}]}]}
    p = tmp_path / 'deck.json'
    p.write_text(json.dumps(spec), encoding='utf-8')
    out = tmp_path / 'd.pptx'
    assert build_deck.main([str(p), '--out', str(out), '--quiet']) == 0
    return zipfile.ZipFile(out).read('ppt/slides/slide1.xml').decode('utf-8')


def test_drawn_shapes_have_no_theme_style(tmp_path):
    xml = _deck(tmp_path)
    for name in ('grat', 'box', 'dot'):
        i = xml.index(f'name="{name}"')
        j = xml.find('</p:sp>', i)
        assert '<p:style>' not in xml[i:j], name
    assert 'Khí áp' in xml


def test_strip_keeps_text_shapes():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
    b.text_frame.text = 'chữ'
    assert build_deck.strip_theme_style(prs) == 1
    assert a._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style') is None
    assert b._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style') is not None


def test_skill_md_creative_authority():
    t = (SKILL / 'SKILL.md').read_text(encoding='utf-8')
    assert 'Version 5.7.0' in t
    sec = t[t.index('## Full creative authority'):t.index('## The one idea')]
    for must in ('full creative authority', 'deck-wide concept', 'core, not an option', '≥ 8', 'floor, not a ceiling',
                 'never', 'flagship-khi-ap'):
        assert must in sec, must


def test_manimgx_quirks_reference():
    t = (SKILL / 'references' / 'manimgx-quirks.md').read_text(encoding='utf-8')
    for must in ('nameID 1', 'matplotlib', 'VFR', 'Stroke widths', 're-adds', 'UpdateFromFunc', 'effectRef', 'strip_theme_style'):
        assert must in t, must


def test_flagship_example_is_light_and_complete():
    ex = SKILL / 'examples' / 'flagship-khi-ap'
    for f in ('README.md', 'art-direction.md', 'critique-log.md', 'films/kit.py', 'films/pressure.py', 'films/baro.py',
              'films/drum.py', 'films/wind.py', 'films/tide.py', 'anatomy/chipline.mjs', 'make_deck.py', 'contact.jpg'):
        assert (ex / f).exists(), f
    size = sum(p.stat().st_size for p in ex.rglob('*') if p.is_file())
    assert size < 6_000_000, size
    assert not list(ex.rglob('*.mp4')) and not list(ex.rglob('*.pptx'))


def test_anatomy_film_names_inch_anchors():
    t = (SKILL / 'scripts' / 'anatomy_film.py').read_text(encoding='utf-8')
    assert 'anchors_in' in t and 'figure pixels' in t
